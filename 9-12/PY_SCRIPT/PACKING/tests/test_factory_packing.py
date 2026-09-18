from copy import deepcopy
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch
import unittest

from openpyxl import Workbook, load_workbook

from config import PROJECT_DIR
from models.packing import OrderItem, PackingDataError
from services.factory_excel_service import write_factory_workbook, FACTORY_REMINDER, SHIPPING_FIELDS
from services.factory_packing_service import (
    BookingLookup, FactoryRow, FactoryScheduleLookup, collect_factory_rows, slip_cartons,
)
from services.product_catalog_service import read_product_catalog
from services.schedule_service import OrderSchedule, read_order_schedule


def item(base='28483', quantity=1800, pack=1):
    return OrderItem(f'{base}-000-1B-001-OPS', Decimal(quantity), Decimal(pack),
                     '001234', 'Playskool Bobble Ball', '', 'BOX WITH SLIP SHEET',
                     customer_item_no='000567')


def record():
    return {'sheet': 'JP', 'row': 5, 'values': {
        '产品名称': '振动球', '装箱数量': 1, '长': 16.83, '宽': 17.15, '高': 20,
        '整箱净重kg': 0.556, '整箱毛重kg': 0.716,
        '整个卡板净重KG': 100, '整个卡板毛重KG': 130,
    }}


class FactoryPackingTests(unittest.TestCase):
    def test_schedule_customers_do_not_affect_packaging_lookup(self):
        lookup = FactoryScheduleLookup(OrderSchedule(Path('s.xlsx'), {'#28483': [
            ('生产订单号', '长编号', '客户', '包装要求'),
            ('POHK-26-13545', item().item_no, 'WM US', '包装 A'),
            ('POHK-26-13545', item().item_no, 'WM CA', '包装 A'),
            ('POHK-26-00001', item().item_no, '其他客户', '其他包装'),
        ]}))
        choose = Mock()
        self.assertEqual(lookup.packaging('POHK-26-13545-001', item(), choose), '包装 A')
        choose.assert_not_called()
        missing = FactoryScheduleLookup(OrderSchedule(Path('s.xlsx'), {'#28483': [
            ('生产订单号', '客户', '包装要求'), ('POHK-26-13545', None, '包装 A')]}))
        self.assertEqual(missing.packaging('POHK-26-13545-001', item(), choose), '包装 A')

    def test_slips_are_matched_per_item_and_require_integral_pallets(self):
        remarks = ('28483 - 180pcs = 180ctns = 1 slip sheet\n'
                   '28481 - 480pcs = 120ctns = 1 slip sheet\n'
                   '28472 - 60pcs = 60ctns = 1 slip sheet')
        self.assertEqual(slip_cartons(item(), remarks), 180)
        self.assertEqual(slip_cartons(item('28481', 1920, 4), remarks), 120)
        with self.assertRaises(PackingDataError):
            slip_cartons(item(quantity=181), remarks)
        with self.assertRaises(PackingDataError):
            slip_cartons(item('99999'), remarks)
        with self.assertRaises(PackingDataError):
            slip_cartons(item(), remarks + '\n28483 - 100pcs = 100ctns = 1 slip sheet')
        loose = item()
        loose.packaging = 'OPEN PLATFORM BOX'
        self.assertIsNone(slip_cartons(loose, remarks))

    def test_schedule_uses_order_and_full_item_then_asks_for_conflicts(self):
        rows = [('生产订单号', '长编号', '包装要求'),
                ('POHK-26-13545', item().item_no, '包装 A'),
                ('POHK-26-13545', item().item_no, '包装 B'),
                ('POHK-26-13545', '28483-000-2B-001-OPS', '不应匹配'),
                ('POHK-26-11111', item().item_no, '其他订单')]
        lookup = FactoryScheduleLookup(OrderSchedule(Path('s.xlsx'), {'#28483': rows}))
        choose = Mock(return_value=1)
        self.assertEqual(lookup.packaging('POHK-26-13545-001', item(), choose), '包装 B')
        self.assertEqual(len(choose.call_args.args[2]), 2)
        with self.assertRaises(PackingDataError):
            lookup.packaging('POHK-26-13545-001', item(), lambda *args: None)
        with self.assertRaises(PackingDataError):
            lookup.packaging('POHK-26-00000-001', item(), choose)

    def test_exact_order_preferred_and_empty_item_can_use_sheet_number(self):
        lookup = FactoryScheduleLookup(OrderSchedule(Path('s.xlsx'), {'#28483 产品': [
            ('生产订单号', '包装要求'), ('POHK-26-13545', '旧包装'),
            ('POHK-26-13545-001', '新包装')]}))
        self.assertEqual(lookup.packaging('POHK-26-13545-001', item(), Mock()), '新包装')

    def test_booking_filename_exact_match_and_conflict_selection(self):
        with TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'sub').mkdir()
            for name in ('001234.pdf', 'sub/001234.PDF', '1001234.pdf'):
                (root / name).touch()
            reader = Mock()
            reader.pages = [Mock(extract_text=lambda: 'Shipper booking number\nSBK0004632022')]
            other = Mock()
            other.pages = [Mock(extract_text=lambda: 'Shipper booking number: SBK0004632023')]
            choose, log = Mock(return_value=1), Mock()
            with patch('services.factory_packing_service.PdfReader', side_effect=[reader, other]) as read:
                lookup = BookingLookup(root, choose, log)
                self.assertEqual(lookup.get('001234'), 'SBK0004632023')
                self.assertEqual(lookup.get('001234'), 'SBK0004632023')
                self.assertEqual(read.call_count, 2)
                self.assertEqual(lookup.get('missing'), '')
                log.assert_called_once()

    def test_known_unnamed_jp_column_is_supported(self):
        with TemporaryDirectory() as temp:
            path = Path(temp) / 'jp.xlsx'
            book = Workbook()
            book.active.append(['序号', '图示', None, '产品名称', '装箱数量', '长', '宽', '高'])
            book.active.append([1, None, 28483, '振动球', 1, 10, 20, 30])
            book.save(path)
            self.assertEqual(read_product_catalog(path).items['28483'][0]['values']['产品名称'], '振动球')

    def test_excel_multi_order_formulas_totals_and_no_sample_residue(self):
        rows = [FactoryRow('POHK-26-13545-001', 'Customer A', item(), record(), '包装', 'SBK01', Decimal(180)),
                FactoryRow('POHK-26-13546-001', 'Customer B', item('28481', 1920, 4), record(), '包装 B', '', Decimal(120))]
        # More rows than the four template rows tests expansion and footer placement.
        rows += [deepcopy(rows[0]) for _ in range(4)]
        with TemporaryDirectory() as temp:
            path = write_factory_workbook(rows, PROJECT_DIR / 'templates/factory_notice.xlsx', temp)
            self.assertEqual(path.name, '9480.xlsx')
            book = load_workbook(path)
            sheet = book.active
            self.assertEqual(book.sheetnames, ['工厂装箱单'])
            self.assertEqual(sheet['C4'].value, '001234')
            self.assertEqual(sheet['E4'].value, '000567')
            self.assertEqual(sheet['N4'].value, 0.556)  # Never pallet weight.
            self.assertEqual(sheet['J4'].value, '=H4/I4')
            self.assertEqual(sheet['R9'].value, '=J9*K9*L9*M9/1000000')
            self.assertEqual(sheet['S5'].value, '=J5/120')
            self.assertEqual(sheet['P5'].value, '=J5*N5')
            self.assertEqual(sheet['J10'].value, '=SUM(J4:J9)')
            self.assertEqual(sheet['B5'].value, 'POHK-26-13546-001')
            self.assertEqual(sheet['A4'].value, 'Customer A')
            self.assertFalse(sheet.column_dimensions['A'].hidden)
            self.assertEqual(sheet.sheet_view.topLeftCell, 'A1')
            self.assertEqual(sheet.sheet_view.pane.topLeftCell, 'A4')
            self.assertIsNone(sheet.sheet_view.pane.xSplit)
            self.assertEqual(sheet.sheet_view.pane.ySplit, 3)
            self.assertEqual(sheet['F13'].value, FACTORY_REMINDER)
            self.assertEqual(sheet['F13'].fill.fgColor.rgb, '00FFFF00')
            self.assertTrue(sheet['F13'].font.bold)
            self.assertEqual(sheet['F13'].font.color.rgb, '00FF0000')
            self.assertEqual([sheet.cell(r, 16).value for r in range(12, 20)], list(SHIPPING_FIELDS))
            self.assertTrue(all(sheet.cell(r, 18).value is None for r in range(12, 20)))
            self.assertIn('$A$1:$Y$19', str(sheet.print_area))
            self.assertIsNone(sheet['V4'].value)
            self.assertIsNone(sheet['W4'].value)
            self.assertIsNone(sheet['X4'].value)
            self.assertTrue(book.calculation.fullCalcOnLoad)
            self.assertNotIn('Y5:Y6', {str(r) for r in sheet.merged_cells.ranges})
            book.close()
            original = path.read_bytes()
            with self.assertRaises(FileExistsError):
                write_factory_workbook(rows, PROJECT_DIR / 'templates/factory_notice.xlsx', temp)
            self.assertEqual(path.read_bytes(), original)

    def test_multiple_pdfs_merge_and_duplicate_orders_abort(self):
        sample = PROJECT_DIR / '工厂箱单'
        if not (sample / 'po_list').is_dir():
            self.skipTest('Local sample unavailable')
        catalog = read_product_catalog(sample / 'JP产品净重毛重外箱尺寸表20241209.xlsx')
        schedule = read_order_schedule(sample / '2026年 可图雅JP订单分类排期汇总（9-12更新).xlsx')
        pdf = next((sample / 'po_list').glob('*.pdf'))
        booking_pdf = next(sample.rglob('1003449527.pdf'))
        choose = lambda order_item, candidates, hint: (candidates[0], order_item.description)
        with patch('services.factory_packing_service.scan_pdfs', side_effect=[[pdf, pdf], [booking_pdf]]), \
             patch('services.factory_packing_service.read_order_number', side_effect=['POHK-26-13545-001', 'POHK-26-13545-002']), \
             patch('services.factory_packing_service.read_order_items', wraps=__import__(
                 'services.order_pdf_service', fromlist=['read_order_items']).read_order_items) as read:
            # Both input files supply real parsed sample rows; vary the PO for the second input.
            from services.order_pdf_service import read_order_items
            read.side_effect = lambda path, order, **kwargs: read_order_items(path, 'POHK-26-13545-001', **kwargs)
            rows = collect_factory_rows(sample / 'po_list', sample, catalog, schedule, choose, Mock())
            self.assertEqual(len(rows), 4)
            self.assertEqual(sum(r.cartons for r in rows), 4560)
        with patch('services.factory_packing_service.scan_pdfs', side_effect=[[pdf, pdf], [booking_pdf]]):
            with self.assertRaisesRegex(PackingDataError, '重复'):
                collect_factory_rows(sample / 'po_list', sample, catalog, schedule, choose, Mock())

    def test_real_sample_has_2280_cartons_and_10_plus_4_slips(self):
        sample = PROJECT_DIR / '工厂箱单'
        if not (sample / 'po_list').is_dir():
            self.skipTest('Local sample unavailable')
        catalog = read_product_catalog(sample / 'JP产品净重毛重外箱尺寸表20241209.xlsx')
        schedule = read_order_schedule(sample / '2026年 可图雅JP订单分类排期汇总（9-12更新).xlsx')
        choices = []

        def choose(order_item, candidates, hint):
            choices.append((order_item.item_no, hint))
            carton = next(r for r in candidates if '卡板' not in str(r['values']['产品名称']))
            return carton, order_item.description

        rows = collect_factory_rows(sample / 'po_list', sample, catalog, schedule, choose,
                                    lambda *args: self.fail('Unexpected conflict'))
        self.assertEqual(len(choices), 2)
        self.assertEqual([r.cartons for r in rows], [1800, 480])
        self.assertEqual([r.cartons / r.cartons_per_slip for r in rows], [10, 4])
        self.assertTrue(all(r.booking == 'SBK0004632022' for r in rows))
        self.assertEqual([r.customer for r in rows], ['Wal-Mart Store Inc.', 'Wal-Mart Store Inc.'])
        self.assertEqual([r.item.customer_item_no for r in rows], ['678696951', '678696954'])
        with TemporaryDirectory() as temp:
            path = write_factory_workbook(rows, PROJECT_DIR / 'templates/factory_notice.xlsx', temp)
            self.assertEqual(path.name, '2280.xlsx')
            book = load_workbook(path)
            self.assertEqual(book.active['J6'].value, '=SUM(J4:J5)')
            self.assertEqual(book.active['G5'].value, '骰子')
            self.assertEqual(book.active['A4'].value, 'Wal-Mart Store Inc.')
            self.assertEqual(book.active['F9'].value, FACTORY_REMINDER)
            self.assertEqual(book.active['P15'].value, '车牌：')
            book.close()


if __name__ == '__main__':
    unittest.main()
