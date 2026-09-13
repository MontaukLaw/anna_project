from datetime import date
from decimal import Decimal
from io import BytesIO
import unittest
from openpyxl import load_workbook
from config import PACKING_TEMPLATE
from models.packing import OrderItem, PackingDataError
from services.slip_sheet_service import build_slip_rows
from services.packing_excel_service import build_packing_excel
from services.packing_display_service import packaging_keywords


class SlipSheetTests(unittest.TestCase):
    def test_two_rows_and_totals_for_multiple_pieces_per_carton(self):
        item = OrderItem('28483', Decimal(3600), Decimal(2), '123', 'Ball', '',
                         packaging='BOX, SLIP SHEET', shipping_remarks='360pcs = 180ctns = 1 slip sheet')
        records = [dict(sheet='JP', row=1, values={'产品名称':'球', '长':10, '宽':20, '高':30,
                    '整箱净重kg':1, '整箱毛重kg':2}),
                   dict(sheet='JP', row=2, values={'产品名称':'球（卡板装）', '长':100, '宽':100, '高':100,
                    '整个卡板净重KG':180, '整个卡板毛重KG':362})]
        rows = build_slip_rows(item, records, 'CODE', date(2026,9,1), '', None, lambda *args:None)
        self.assertEqual(rows[1].display_quantity, 1800)
        self.assertEqual(rows[1].display_case_pack, 180)
        book = load_workbook(BytesIO(build_packing_excel(PACKING_TEMPLATE, 'POHK-26-12345-001', 'Customer', rows)))
        try:
            s = book.active
            self.assertEqual(s['F20'].value, 1800)
            self.assertEqual(s['F21'].value, 10)
            self.assertEqual(s['AA20'].value, 'cargo')
            self.assertEqual(s['AA21'].value, 'cargo together with slip sheet')
            self.assertEqual(s['W21'].value, '=I21*F21')
            self.assertIsNone(s['E22'].value)
            self.assertIn('180 ctns = 1 slip sheet', s['A22'].value)
            self.assertEqual(s['A22'].fill.fgColor.rgb, '00FFF2CC')
            summary = {s.cell(r, 1).value: s.cell(r, 7).value for r in range(24, s.max_row + 1)}
            self.assertEqual(summary['Product weight + Carton Box Weight :'], 360)
            self.assertEqual(summary['Total others packing material weight as 彩盒，说明书，卡纸类:'], 2)
            self.assertEqual(summary['GRAND TOTAL CARTON GROSS WEIGHT PER CARTON :'], 362)
            self.assertEqual(summary['訂單該 ITEM 的 slip sheet 數量'], 10)
            self.assertEqual(summary['每 slip sheet 的 CTN 數量'], 180)
            self.assertEqual(summary['每 slip sheet 的總體積'], 1)
        finally:
            book.close()
        self.assertEqual(packaging_keywords(item.packaging)[0].group(), 'SLIP SHEET')
        item.shipping_remarks = ''
        with self.assertRaises(PackingDataError):
            build_slip_rows(item, records, 'CODE', date(2026,9,1), '', None, lambda *args:None)
