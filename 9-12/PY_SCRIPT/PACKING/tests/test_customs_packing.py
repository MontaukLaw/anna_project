from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook

from services.customs_packing_service import (read_customs_packing, packing_summary, row_message,
                                             result_columns, header_message)


SAMPLES = Path(__file__).resolve().parents[4] / 'asn'


class CustomsPackingTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'packing.xlsx'
        self.book = Workbook()
        self.addCleanup(self.book.close)
        sheet = self.book.active
        sheet.title = '货物'
        sheet.append(['装箱清单'])
        sheet.append(['Customer', 'PO#', 'CUSTOMER ORDER NO.', 'Item#', 'Customer Item No.:',
                      'Description', None, 'Qty/PCS', 'Qty/Ctn', 'no.of ctn',
                      'Measurement/ctn (L/W/H ) cm', None, None, 'NW/CTN', 'GW/CTN', 'NW', 'GW',
                      'CBM', ' CUSTOMER PO #', '品牌', '包装方式', '单价', '表面材质', '出货日期', 'SO'])
        for column in ('A', 'B', 'C', 'D', 'E', 'H', 'I', 'J', 'N', 'O', 'R', 'S', 'T', 'U', 'W', 'X', 'Y'):
            sheet.merge_cells(f'{column}2:{column}3')
        sheet.merge_cells('F2:G3')
        sheet.merge_cells('K2:M3')
        sheet['P3'], sheet['Q3'], sheet['V3'] = 'KGS', 'KGS', 'USD'
        self.row = ['Test Customer', 'PO001', '000123', '0011369', 'CI001', 'Toy', '玩具',
                    12, 4, 3, 42.86, 43.18, 33.34, 2.7, 3.44, 8.1, 10.32, .185106493896,
                    '000782', 'Disney', '1PC/彩盒', 7.5, '毛绒', 46269, 'SO001']
        sheet.append(self.row)
        sheet.append(self.row)  # Identical rows must remain separate ASN source rows.
        sheet['G6'], sheet['H6'], sheet['J6'] = 'Total：', 24, 6
        sheet['B8'] = '24PCS/6箱'
        sheet['P10'], sheet['Q10'] = '出口国家：', '美国'
        self.book.create_sheet('空白')

    def read(self):
        self.book.save(self.path)
        return read_customs_packing(self.path)

    def test_all_columns_dates_source_order_and_duplicate_rows_are_preserved(self):
        data = self.read()
        self.assertEqual(len(data.rows), 2)
        self.assertEqual([row.source_row for row in data.rows], [4, 5])
        self.assertEqual(len(data.rows[0].values), 25)
        self.assertEqual(data.rows[0].values['item'], '0011369')
        self.assertEqual(data.rows[0].values['customer_po'], '000782')
        self.assertEqual(data.rows[0].values['chinese_name'], '玩具')
        self.assertEqual(data.rows[0].values['width'], Decimal('43.18'))
        self.assertEqual(data.rows[0].values['ship_date'], date(2026, 9, 4))
        self.assertEqual(data.rows[0].raw_values['ship_date'], 46269)
        self.assertIn('合计箱数 6', packing_summary(data))
        self.assertIn('合计数量 24', packing_summary(data))
        columns = result_columns(data)
        self.assertEqual(len(header_message(columns).split(' | ')), 26)
        self.assertEqual(len(row_message(1, data.rows[0], columns).split(' | ')), 26)
        self.assertTrue(row_message(1, data.rows[0]).endswith(' | SO001'))
        data.rows[0].values['description'] = 'Toy\nwith | character'
        self.assertNotIn('\n', row_message(1, data.rows[0]))
        self.assertIn('Toy with ｜ character', row_message(1, data.rows[0]))
        self.assertFalse(data.warnings)
        before = self.path.read_bytes()
        read_customs_packing(self.path)
        self.assertEqual(before, self.path.read_bytes())

    def test_blank_rows_skipped_but_zero_and_incomplete_rows_kept(self):
        sheet = self.book.active
        sheet.insert_rows(5)
        sheet['H6'], sheet['J6'], sheet['D6'] = 0, 0, None
        data = self.read()
        self.assertEqual([row.source_row for row in data.rows], [4, 6])
        self.assertEqual(data.rows[1].values['quantity'], 0)
        self.assertIsNone(data.rows[1].values['item'])
        self.assertTrue(any('货号为空' in warning for warning in data.warnings))

    def test_formula_missing_cache_and_excel_error_are_visible_without_dropping_rows(self):
        sheet = self.book.active
        sheet['H4'] = '=I4*J4'
        sheet['P5'] = '#VALUE!'
        data = self.read()
        self.assertEqual(len(data.rows), 2)
        self.assertIsNone(data.rows[0].values['quantity'])
        self.assertEqual(data.rows[0].raw_values['quantity'], '=I4*J4')
        self.assertTrue(any('没有缓存值' in warning for warning in data.warnings))
        self.assertTrue(any('Excel 错误' in warning for warning in data.warnings))
        self.assertNotIn('合计数量', packing_summary(data))

    def test_multiple_sheets_keep_workbook_order_and_formatted_identifiers(self):
        sheet = self.book.active
        sheet['D4'], sheet['D4'].number_format = 11369, '0000000'
        self.book.copy_worksheet(sheet).title = '第二批'
        data = self.read()
        self.assertEqual([row.sheet for row in data.rows], ['货物', '货物', '第二批', '第二批'])
        self.assertEqual(data.rows[0].values['item'], '0011369')

    def test_wrong_file_and_unrecognized_headers_fail(self):
        with self.assertRaises(ValueError):
            read_customs_packing(self.path.with_suffix('.xls'))
        self.book.active['H2'] = 'unknown'
        with self.assertRaisesRegex(ValueError, '未找到海关装箱单明细'):
            self.read()

    def test_merged_so_and_shared_text_are_read_for_each_detail(self):
        sheet = self.book.active
        sheet['Y4'] = 'SBK0004178305'
        for column in ('A', 'B', 'Y'):
            sheet.merge_cells(f'{column}4:{column}5')
        data = self.read()
        self.assertEqual([row.source_row for row in data.rows], [4, 5])
        self.assertEqual([row.values['so'] for row in data.rows], ['SBK0004178305'] * 2)
        self.assertEqual([row.values['customer'] for row in data.rows], ['Test Customer'] * 2)
        self.assertEqual([row.values['po'] for row in data.rows], ['PO001'] * 2)
        self.assertTrue(row_message(2, data.rows[1]).endswith(' | SBK0004178305'))
        self.assertIsNone(data.rows[1].raw_values['so'])  # Preserve the physical source cell.
        self.assertFalse(data.warnings)
        before = self.path.read_bytes()
        read_customs_packing(self.path)
        self.assertEqual(before, self.path.read_bytes())

    def test_merged_so_respects_group_boundaries_and_skips_blank_rows(self):
        sheet = self.book.active
        sheet.delete_rows(6, 5)  # Replace fixture footer with a second group and total.
        sheet.append(self.row)
        sheet.append(self.row)
        sheet.append(self.row)
        sheet['Y4'] = 'SBK0004178305'
        sheet['Y6'] = 123
        sheet['Y6'].number_format = '000000'
        sheet['Y8'] = None
        sheet.merge_cells('Y4:Y5')
        sheet.merge_cells('Y6:Y7')
        sheet.merge_cells('Y8:Y9')
        sheet['G10'] = 'Total'
        data = self.read()
        self.assertEqual([row.source_row for row in data.rows], [4, 5, 6, 7, 8])
        self.assertEqual([row.values['so'] for row in data.rows],
                         ['SBK0004178305', 'SBK0004178305', '000123', '000123', None])

    def test_merged_so_reaches_both_warehouse_asn_rows(self):
        from config import ASN_TEMPLATE, ZHONGTONG_ASN_TEMPLATE
        from services.asn_generation_service import build_asn_plan, generation_options
        from services.hs_units_service import LegalUnits
        from test_asn_generation import fixture

        self.book.active['Y4'] = 'SBK0004178305'
        self.book.active.merge_cells('Y4:Y5')
        packing = self.read()
        report, _ = fixture()
        for document in report.documents['TEST001'].values():
            item = document.items[0]
            item.name = '填充玩具/玩具'
            item.quantity, item.amount = Decimal(24), Decimal(180)
            item.cartons = Decimal(6)
            item.net_weight, item.gross_weight = Decimal('16.2'), Decimal('20.64')
        for warehouse, template in [('yixing', ASN_TEMPLATE), ('zhongtong', ZHONGTONG_ASN_TEMPLATE)]:
            with self.subTest(warehouse=warehouse):
                options = dict(generation_options(report, packing), warehouse=warehouse,
                               legal_units={'9503008390': LegalUnits('个', '千克')})
                plan = build_asn_plan(report, packing, template, options)
                self.assertEqual([row[1] for row in plan['rows']], ['SBK0004178305'] * 2)

    def test_merged_formulas_warn_and_numeric_totals_are_not_duplicated(self):
        sheet = self.book.active
        sheet['Y4'] = '="SBK0004178305"'
        sheet.merge_cells('Y4:Y5')
        sheet.merge_cells('H4:H5')
        data = self.read()
        self.assertEqual(len(data.rows), 2)
        self.assertIsNone(data.rows[1].values['quantity'])
        self.assertTrue(any('第 5 行：SO 的公式没有缓存值' in w for w in data.warnings))
        self.assertTrue(any('第 5 行：数量为空' in w for w in data.warnings))

    @unittest.skipUnless(SAMPLES.exists(), 'Local samples unavailable')
    def test_supplied_sample_has_twelve_rows_and_309_cartons(self):
        paths = list(SAMPLES.glob('PL-*.xlsx'))
        if not paths:
            self.skipTest('Local PL sample unavailable')
        data = read_customs_packing(paths[0])
        self.assertEqual(len(data.rows), 12)
        self.assertEqual([row.source_row for row in data.rows], list(range(4, 16)))
        self.assertEqual(sum(row.values['cartons'] for row in data.rows), 309)
        self.assertEqual(sum(row.values['quantity'] for row in data.rows), 1236)
        self.assertEqual(data.rows[0].values['so'], 'ZLC26-2192')
        self.assertEqual(data.rows[-1].values['so'], 'ZLC26-2384')
        self.assertEqual(len(data.rows[0].values), 25)
        self.assertFalse(data.warnings)
