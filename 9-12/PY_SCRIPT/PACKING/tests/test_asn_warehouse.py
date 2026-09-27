from copy import copy
from datetime import date, datetime
import os
from pathlib import Path
import subprocess
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from config import ZHONGTONG_ASN_TEMPLATE
from services.asn_generation_service import output_filename, write_asn_workbook
from services.asn_warehouse import warehouse_profile


class WarehouseTests(unittest.TestCase):
    def test_distinct_output_names_and_unknown_warehouse_rejected(self):
        self.assertTrue(output_filename('TEST001').endswith('TEST001.xls'))
        self.assertEqual(output_filename('TEST001', 'zhongtong'), '中通仓-入仓预申报ASN表格-TEST001.xlsx')
        with self.assertRaises(ValueError):
            warehouse_profile('unknown')
        with self.assertRaises(ValueError):
            output_filename('../bad', 'zhongtong')

    @unittest.skipUnless(os.name == 'nt', 'Requires Windows and Microsoft Excel')
    def test_native_zhongtong_writer_preserves_template_and_calculates(self):
        original = ZHONGTONG_ASN_TEMPLATE.read_bytes()
        row = [None] * 30
        for index, value in {0: 1, 1: 'SO001', 2: '000123-PO', 3: '00123',
                             4: 4, 5: 4, 6: 1, 7: 3.44, 8: 2.7, 12: 4, 14: 7.5, 23: 51169}.items():
            row[index] = value
        plan = {'warehouse': 'zhongtong', 'template': str(ZHONGTONG_ASN_TEMPLATE),
                'sheet': 'ASN', 'scope': 'all', 'contract': 'TEST001',
                'headers': {'D3': date.today().isoformat(), 'D6': 'Test customer', 'H11': 'TEST001',
                            'H12': '2026-05-20', 'H24': 3},
                'rows': [row, list(row), list(row)],
                'formulas': [{'cell': f'P{r}', 'value': f'=ROUND(M{r}*O{r},2)'} for r in range(26, 29)]}
        with TemporaryDirectory() as folder:
            output = write_asn_workbook(plan, folder)
            template = load_workbook(ZHONGTONG_ASN_TEMPLATE)
            book = load_workbook(output)
            cached = load_workbook(output, data_only=True)
            try:
                self.assertEqual(book.sheetnames, template.sheetnames)
                sheet = book['ASN']
                self.assertEqual(sheet['D6'].value, 'Test customer')
                self.assertEqual(sheet['D26'].value, '00123')
                self.assertEqual(sheet['C26'].value, '000123-PO')
                self.assertEqual(sheet['H24'].value, 3)
                self.assertEqual(sheet['H24'].data_type, 'n')
                self.assertIsInstance(sheet['D3'].value, datetime)
                self.assertEqual(sheet['D3'].value.date(), date.today())
                self.assertEqual(sheet['H12'].value, datetime(2026, 5, 20))
                self.assertEqual(sheet['X26'].value, 51169)
                self.assertEqual(sheet['X26'].data_type, 'n')
                self.assertEqual(cached['境内货源地代码']['E1'].value, '达县')
                self.assertEqual(cached['境内货源地代码']['F1'].value, '51169达县')
                self.assertEqual(sheet['P26'].value, '=ROUND(M26*O26,2)')
                self.assertEqual(cached['ASN']['P26'].value, 30)
                self.assertTrue(all(c.value is None for row in sheet.iter_rows(min_row=29, max_col=30)
                                    for c in row))
                self.assertEqual(sheet.print_area, template['ASN'].print_area)
                for name in template.sheetnames:
                    self.assertEqual(book[name].sheet_state, template[name].sheet_state, name)
                    self.assertEqual(set(map(str, book[name].merged_cells.ranges)),
                                     set(map(str, template[name].merged_cells.ranges)), name)
                    self.assertEqual(book[name].page_setup.scale, template[name].page_setup.scale, name)
                for row_number in range(26, 29):
                    self.assertEqual(sheet.row_dimensions[row_number].height,
                                     template['ASN'].row_dimensions[row_number].height)
                    for column in range(1, 31):
                        before = template['ASN'].cell(row_number, column)
                        after = sheet.cell(row_number, column)
                        for component in ('font', 'fill', 'border', 'alignment', 'number_format', 'protection'):
                            self.assertEqual(copy(getattr(before, component)), copy(getattr(after, component)),
                                             (before.coordinate, component))
            finally:
                template.close()
                book.close()
                cached.close()
            comparison = subprocess.run(
                ['powershell.exe', '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass',
                 '-File', str(Path(__file__).with_name('compare_asn_template.ps1')),
                 '-TemplatePath', str(ZHONGTONG_ASN_TEMPLATE), '-OutputPath', str(output),
                 '-AsnSheetName', 'ASN'], capture_output=True, text=True, encoding='utf-8',
                errors='replace', creationflags=subprocess.CREATE_NO_WINDOW, timeout=120)
            self.assertEqual(comparison.returncode, 0, comparison.stdout + comparison.stderr)
            self.assertIn('PRESERVED_SHEETS=17', comparison.stdout)
        self.assertEqual(ZHONGTONG_ASN_TEMPLATE.read_bytes(), original)
