from copy import copy
from datetime import datetime
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook

from config import XINGHUI_ASN_TEMPLATE
from services.asn_generation_service import output_filename, write_asn_workbook, build_asn_plan, generation_options
from services.asn_warehouse import warehouse_profile
from services.hs_units_service import LegalUnits
from services.asn_pallet_service import pallet_counts


class XinghuiTests(unittest.TestCase):
    def sources(self):
        from test_asn_generation import fixture
        report, packing = fixture()
        report.documents['TEST001']['香港合同'].items[0].specification = '3|0|玩具|Playskool牌/Hasbro牌/JUST PLAY牌'
        for row in packing.rows:
            row.values['brand'] = 'PL-only brand'
            row.values['packaging'] = '4PC/箱，1箱/卡板+滑片纸'
        options = dict(generation_options(report, packing), warehouse='xinghui',
                       legal_units={'9503008390': LegalUnits('套', '千克')})
        return report, packing, options

    def test_confirmed_fields_and_contract_brand(self):
        report, packing, options = self.sources()
        plan = build_asn_plan(report, packing, XINGHUI_ASN_TEMPLATE, options)
        self.assertEqual(plan['headers']['G8'], 'Just Play LLC')
        self.assertEqual(plan['headers']['H15'], 'TEST001')
        self.assertIsNone(plan['headers']['H14'])
        self.assertIsNone(plan['headers']['F19'])
        row = plan['rows'][0]
        self.assertEqual(row[3:5], ['000123-PO', '000987-SKU'])
        self.assertEqual(row[8:14], [4, '套', 4, 1, 3.44, 2.7])
        self.assertEqual(row[17:20], [4, '套', 7.5])
        self.assertEqual(row[22:27], [4, '套', 2.7, '千克', 'Playskool牌/Hasbro牌/JUST PLAY牌'])
        self.assertEqual(row[31:36], [40, 30, 20, 'CTN', .024])
        self.assertEqual(row[44:46], ['达县', None])
        self.assertEqual(plan['formulas'][0], {'cell': 'U26', 'value': '=ROUND(R26*T26,2)'})

    def test_missing_contract_brand_is_not_replaced_by_pl_brand(self):
        report, packing, options = self.sources()
        for spec in ('3|0|玩具', '3|0|玩具|Disney牌|Hasbro牌'):
            report.documents['TEST001']['香港合同'].items[0].specification = spec
            with self.subTest(spec=spec), self.assertRaisesRegex(ValueError, '合同品牌'):
                build_asn_plan(report, packing, XINGHUI_ASN_TEMPLATE, options)

    def test_template_row_limit_and_customer_item_required(self):
        report, packing, options = self.sources()
        packing.rows[0].values['customer_item'] = ''
        with self.assertRaisesRegex(ValueError, 'customer_item 未填写'):
            build_asn_plan(report, packing, XINGHUI_ASN_TEMPLATE, options)
        packing.rows = [packing.rows[0]] * 972
        with self.assertRaisesRegex(ValueError, '最多可填写 971 行'):
            build_asn_plan(report, packing, XINGHUI_ASN_TEMPLATE, options)

    def test_pallet_fields_and_missing_packaging(self):
        report, packing, options = self.sources()
        plan = build_asn_plan(report, packing, XINGHUI_ASN_TEMPLATE, options)
        self.assertEqual(plan['headers']['D16'], '是')
        self.assertEqual([row[29:31] for row in plan['rows']], [[1, 1], [1, 1]])
        self.assertTrue(all(row[34] == 'CTN' for row in plan['rows']))
        packing.rows[0].values['packaging'] = '4PC/箱，需要卡板'
        with self.assertRaisesRegex(ValueError, '每板箱数缺失或不唯一'):
            build_asn_plan(report, packing, XINGHUI_ASN_TEMPLATE, options)

    def test_pallet_conversion_is_explicit_and_integral(self):
        from decimal import Decimal
        _, packing, _ = self.sources()
        row = packing.rows[0]
        row.values.update(cartons=Decimal(1800), packaging='1PC/彩盒/箱,180箱/卡板+滑片纸')
        self.assertEqual(pallet_counts(row), (10, 180))
        row.values.update(cartons=Decimal(480), packaging='12PCS为1套/彩盒，4盒/箱，120箱/卡板+滑片纸（包3边）')
        self.assertEqual(pallet_counts(row), (4, 120))
        for text in ('需要卡板', '120箱/卡板 或 60箱/卡板', '0箱/卡板', '119箱/卡板'):
            row.values['packaging'] = text
            with self.subTest(text=text), self.assertRaises(ValueError):
                pallet_counts(row)

    def test_cartons_and_pallets_can_share_one_asn(self):
        report, packing, options = self.sources()
        packing.rows[0].values['packaging'] = '2PC/胶袋，12袋/PDQ/箱'
        plan = build_asn_plan(report, packing, XINGHUI_ASN_TEMPLATE, options)
        self.assertEqual(plan['headers']['D16'], '是')
        self.assertEqual([row[29:31] for row in plan['rows']], [[None, None], [1, 1]])
        self.assertTrue(all(row[34] == 'CTN' for row in plan['rows']))
        # The header must also work when the last row is loose cartons.
        packing.rows.reverse()
        plan = build_asn_plan(report, packing, XINGHUI_ASN_TEMPLATE, options)
        self.assertEqual(plan['headers']['D16'], '是')
        self.assertEqual([row[29:31] for row in plan['rows']], [[1, 1], [None, None]])
        for row in packing.rows:
            row.values['packaging'] = '2PC/胶袋，12袋/PDQ/箱'
        plan = build_asn_plan(report, packing, XINGHUI_ASN_TEMPLATE, options)
        self.assertEqual(plan['headers']['D16'], '否')
        self.assertTrue(all(row[29:31] == [None, None] for row in plan['rows']))

    def test_reported_three_rows_and_missing_packaging(self):
        from copy import deepcopy
        from decimal import Decimal
        _, packing, _ = self.sources()
        row = deepcopy(packing.rows[0])
        results = []
        for item, cartons, packaging in [
            ('33465', 180, '2PC/胶袋，12袋/PDQ/箱'),
            ('33465', 180, '2PC/胶袋，12袋/PDQ/箱'),
            ('28481', 480, '12PCS为1套/彩盒，4盒/箱，120箱/卡板+滑片纸（包3边）'),
        ]:
            row.values.update(item=item, cartons=Decimal(cartons), packaging=packaging)
            results.append(pallet_counts(row))
        self.assertEqual(results, [(None, None), (None, None), (4, 120)])
        for packaging in (None, '', ' ', 'BOX WITH SLIP SHEET', 'PALLET', '需托盘'):
            row.values['packaging'] = packaging
            with self.subTest(packaging=packaging), self.assertRaises(ValueError):
                pallet_counts(row)

    def test_profile_and_macro_enabled_filename(self):
        self.assertEqual(output_filename('TEST001', 'xinghui'), '星辉仓-入仓预申报ASN表格-TEST001.xlsm')
        profile = warehouse_profile('xinghui')
        self.assertEqual(profile.file_format, 52)
        self.assertEqual(profile.data_last_row, 996)

    @unittest.skipUnless(os.name == 'nt', 'Requires Windows and Microsoft Excel')
    def test_native_macro_protection_and_template_preserved(self):
        original = XINGHUI_ASN_TEMPLATE.read_bytes()
        row = [None] * 47
        for index, value in {0: 1, 1: 'SO001', 3: '000123-PO', 4: '00123', 8: 4, 9: '套',
                             10: 4, 11: 1, 12: 3.44, 13: 2.7, 17: 4, 18: '套', 19: 7.5}.items():
            row[index] = value
        plan = {'warehouse': 'xinghui', 'template': str(XINGHUI_ASN_TEMPLATE),
                'sheet': '报关资料与ASN', 'scope': 'all', 'contract': 'TEST001',
                'headers': {'D5': '2026-09-26', 'G8': 'Test customer', 'H11': 'TEST001', 'H12': '2026-06-20'},
                'rows': [row], 'formulas': [{'cell': 'U26', 'value': '=ROUND(R26*T26,2)'}]}
        with TemporaryDirectory() as folder:
            output = write_asn_workbook(plan, folder)
            with ZipFile(XINGHUI_ASN_TEMPLATE) as source, ZipFile(output) as exported:
                self.assertEqual(exported.read('xl/vbaProject.bin'), source.read('xl/vbaProject.bin'))
                self.assertIn(b'macroEnabled', exported.read('[Content_Types].xml'))
            template = load_workbook(XINGHUI_ASN_TEMPLATE, keep_vba=True)
            book = load_workbook(output, keep_vba=True)
            cached = load_workbook(output, data_only=True)
            try:
                self.assertEqual(book.sheetnames, template.sheetnames)
                sheet = book['报关资料与ASN']
                self.assertEqual(sheet['G8'].value, 'Test customer')
                self.assertEqual(sheet['D5'].value, datetime(2026, 9, 26))
                self.assertEqual(sheet['H12'].value, datetime(2026, 6, 20))
                self.assertEqual(sheet['D26'].value, '000123-PO')
                self.assertEqual(sheet['E26'].value, '00123')
                self.assertEqual(cached[sheet.title]['U26'].value, 30)
                self.assertTrue(all(sheet.cell(27, col).value is None for col in range(1, 48)))
                for name in template.sheetnames:
                    self.assertEqual(book[name].sheet_state, template[name].sheet_state, name)
                    self.assertEqual(book[name].protection, template[name].protection, name)
                    self.assertEqual(book[name].sheet_properties.codeName,
                                     template[name].sheet_properties.codeName, name)
                    self.assertEqual(book[name].page_setup.scale, template[name].page_setup.scale, name)
                    self.assertEqual(book[name].print_area, template[name].print_area, name)
                for col in range(1, 48):
                    before, after = template[sheet.title].cell(26, col), sheet.cell(26, col)
                    for component in ('font', 'fill', 'border', 'alignment', 'number_format', 'protection'):
                        self.assertEqual(copy(getattr(before, component)), copy(getattr(after, component)),
                                         (before.coordinate, component))
                for pos in ('AX329', 'AY125', 'BA27'):
                    self.assertEqual(sheet[pos].value, template[sheet.title][pos].value)
            finally:
                template.close()
                book.close()
                cached.close()
        self.assertEqual(XINGHUI_ASN_TEMPLATE.read_bytes(), original)
