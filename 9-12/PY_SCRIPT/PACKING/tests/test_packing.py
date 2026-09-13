from decimal import Decimal
from datetime import date
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import unittest
from openpyxl import load_workbook

from config import PACKING_TEMPLATE, PROJECT_DIR, WORKSPACE_DIR
from models.packing import OrderItem, PackingRow
from services.order_pdf_service import read_order_items
from services.packing_excel_service import build_packing_excel
from services.packing_generation_service import generate_packing_lists
from services.product_catalog_service import read_product_catalog
from services.schedule_service import read_order_schedule
from services.settings_service import find_local_product_files, find_local_schedule_files


class PackingTests(unittest.TestCase):
    def test_calculations_and_dynamic_template_rows(self):
        item = OrderItem('12345-000-1A-012-ABC', Decimal(120), Decimal(12), '000123', 'Duck', 'Sep 22, 2026')
        row = PackingRow(item, '2466KTY01', Decimal('2'), Decimal('3'), Decimal('10'), Decimal('20'), Decimal('30'), 'Toy duck', 'JP row 6', date(2026,9,17))
        for count in (1, 4, 7):
            content = build_packing_excel(PACKING_TEMPLATE, 'POHK-26-12345-001', 'Customer', [row]*count)
            book = load_workbook(BytesIO(content), data_only=True)
            try:
                sheet = book.active
                self.assertEqual(sheet['B20'].value, '000123')
                self.assertEqual(sheet['F20'].value, 10)
                self.assertAlmostEqual(sheet['Q20'].value, .006)
                self.assertAlmostEqual(sheet['S20'].value, .060)
                self.assertEqual(sheet['U20'].value, 20)
                self.assertEqual(sheet['W20'].value, 30)
                self.assertEqual(sheet.cell(20+count,5).value,120*count)
                self.assertEqual(sheet.cell(20+count,6).value,10*count)
                self.assertEqual(sheet.cell(22+count,9).value,'91511700MAACP9B4X9')
                self.assertEqual(sheet['T14'].value,'2466KTY01')
                self.assertEqual(sheet['Z20'].value,'N')
                self.assertEqual(sheet['C12'].value.date(),date(2026,9,17))
                self.assertEqual(sheet['C12'].number_format,'yyyy-mm-dd')
                self.assertFalse(sheet._images)
                self.assertTrue(all(not cell.data_type=='f' for values in sheet for cell in values))
            finally:
                book.close()

    def test_existing_target_skips_before_pdf_read(self):
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root/'POHK-26-12345-001_装箱单.xlsx'
            path.write_bytes(b'existing')
            with patch('services.packing_generation_service.read_order_items') as read:
                generated, failed, skipped = generate_packing_lists({'POHK-26-12345-001':[]},None,None,PACKING_TEMPLATE,root,None,lambda *args:None)
                self.assertEqual((generated, failed, skipped), ([], [], ['POHK-26-12345-001']))
                read.assert_not_called()
            self.assertEqual(path.read_bytes(),b'existing')


class LocalPackingIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        products = find_local_product_files(PROJECT_DIR)
        schedules = find_local_schedule_files(PROJECT_DIR)
        if not products or not schedules or not list(WORKSPACE_DIR.glob('*.pdf')):
            raise unittest.SkipTest('Local sample documents not available')
        cls.catalog=read_product_catalog(products[0])
        cls.schedule=read_order_schedule(schedules[0])

    def test_six_real_orders_in_temporary_directory(self):
        import re
        from services.toy_translation_service import translate_toy_name
        matches={re.search(r'POHK-\d+-\d+-\d+',p.name)[0]:[p] for p in WORKSPACE_DIR.glob('*.pdf')}
        choices=[]
        def choose(item, candidates, hint):
            self.assertIn(item.item_no, hint)
            # Test fixture chooses the first catalog row; production always asks the user.
            choices.append(item.item_no)
            return candidates[0],translate_toy_name(candidates[0]['values']['产品名称'])
        with TemporaryDirectory() as temporary:
            generated,failed,skipped=generate_packing_lists(matches,self.catalog,self.schedule,PACKING_TEMPLATE,Path(temporary),choose,lambda *args:None)
            self.assertEqual(failed,[])
            self.assertEqual(len(generated),len(matches))
            self.assertTrue(choices)
            book=load_workbook(Path(temporary)/'POHK-26-16160-001_装箱单.xlsx',data_only=True)
            self.assertEqual(book.active['B20'].value,'0002569483')
            self.assertEqual(book.active['T14'].value,'2366KTY01')
            self.assertEqual(book.active['F20'].value,3420)
            pdf_item=read_order_items(matches['POHK-26-16160-001'][0],'POHK-26-16160-001')[0][0]
            from services.date_service import parse_schedule_date
            self.assertEqual(book.active['C12'].value.date(),parse_schedule_date(pdf_item.cargo_date))
            self.assertEqual(book.active['Y20'].value,pdf_item.description)
            book.close()

    def test_multipage_pdf_includes_all_seven_main_items(self):
        samples=list(WORKSPACE_DIR.parent.glob('*09966*.pdf'))
        if not samples:
            self.skipTest('Multi-item example not available')
        items,_=read_order_items(samples[0],'POHK-26-09966-016')
        self.assertEqual(len(items),7)
        self.assertEqual(items[-1].item_no,'46897-000-1A-004-SWG')
        self.assertEqual(sum(i.quantity for i in items),9004)

    def test_existing_order_does_not_stop_next_order(self):
        from services.toy_translation_service import translate_toy_name
        orders=['POHK-26-13552-001','POHK-26-14888-001']
        matches={order:[next(WORKSPACE_DIR.glob(f'*{order}*.pdf'))] for order in orders}
        def choose(item,candidates,hint):
            return candidates[0],translate_toy_name(candidates[0]['values']['产品名称'])
        with TemporaryDirectory() as temporary:
            root=Path(temporary)
            existing=root/(orders[0]+'_装箱单.xlsx')
            existing.write_bytes(b'keep existing workbook')
            with patch('services.packing_generation_service.read_order_items',wraps=read_order_items) as read:
                generated,failed,skipped=generate_packing_lists(matches,self.catalog,self.schedule,PACKING_TEMPLATE,root,choose,lambda *args:None)
                self.assertEqual(read.call_count,1)
                self.assertEqual(read.call_args.args[1],orders[1])
            self.assertEqual(skipped,[orders[0]])
            self.assertEqual(failed,[])
            self.assertEqual(generated,[root/(orders[1]+'_装箱单.xlsx')])
            self.assertTrue(generated[0].is_file())
            self.assertEqual(existing.read_bytes(),b'keep existing workbook')

    def test_open_after_save_and_open_failure_continues(self):
        orders=['POHK-26-13552-001','POHK-26-14888-001']
        matches={order:[next(WORKSPACE_DIR.glob(f'*{order}*.pdf'))] for order in orders}
        calls=[]
        logs=[]
        with TemporaryDirectory() as temporary:
            root=Path(temporary)
            def opened(path):
                self.assertTrue(path.is_file())
                calls.append(path.stem.removesuffix('_装箱单'))
                if len(calls)==1:
                    self.assertFalse((root/(orders[1]+'_装箱单.xlsx')).exists())
                    raise OSError('Office unavailable')
            generated,failed,skipped=generate_packing_lists(matches,self.catalog,self.schedule,PACKING_TEMPLATE,root,
                lambda item,candidates,hint:(candidates[0],item.description),
                lambda *args:logs.append(args),on_generated=opened)
            self.assertEqual(calls,orders)
            self.assertEqual(len(generated),2)
            self.assertEqual((failed,skipped),([],[]))
            self.assertTrue(any('Office unavailable' in entry[0] for entry in logs))

    def test_missing_date_code_does_not_write_partial_order(self):
        from services.schedule_service import OrderSchedule
        pdf=next(WORKSPACE_DIR.glob('*13552*.pdf'))
        with TemporaryDirectory() as temporary:
            directory=Path(temporary)
            generated,failed,skipped=generate_packing_lists({'POHK-26-13552-001':[pdf]},self.catalog,
                OrderSchedule(Path('empty.xlsx'),{}),PACKING_TEMPLATE,directory,None,lambda *args:None)
            self.assertEqual(generated,[])
            self.assertEqual(failed,['POHK-26-13552-001'])
            self.assertFalse(list(directory.glob('*.xlsx')))


if __name__ == '__main__':
    unittest.main()
