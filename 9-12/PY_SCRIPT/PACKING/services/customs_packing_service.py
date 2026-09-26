"""Read the PL customs packing list, preserving row order, duplicates and missing values."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from models.customs_packing import CustomsColumn, CustomsPackingData, CustomsPackingRow


# The Description and Measurement headings span multiple columns in the sample.
FIELDS = {
    'customer': ('customer', '客户'),
    'po': ('po', '订单号'),
    'customerorderno': ('customer_order', '客户订单号'),
    'item': ('item', '货号'),
    'customeritemno': ('customer_item', '客户货号'),
    'description': ('description', '英文品名'),
    'qtypcs': ('quantity', '数量'),
    'qtyctn': ('per_carton', '每箱数量'),
    'noofctn': ('cartons', '箱数'),
    'measurementctnlwhcm': ('length', '长(cm)'),
    'nwctn': ('net_per_carton', '单箱净重(kg)'),
    'gwctn': ('gross_per_carton', '单箱毛重(kg)'),
    'nw': ('net_weight', '总净重(kg)'),
    'gw': ('gross_weight', '总毛重(kg)'),
    'cbm': ('volume', '体积(CBM)'),
    'customerpo': ('customer_po', '客户PO'),
    '品牌': ('brand', '品牌'),
    '包装方式': ('packaging', '包装方式'),
    '单价': ('price', '单价(USD)'),
    '表面材质': ('material', '表面材质'),
    '出货日期': ('ship_date', '出货日期'),
    'so': ('so', 'SO'),
}
NUMERIC = {'quantity', 'per_carton', 'cartons', 'length', 'width', 'height',
           'net_per_carton', 'gross_per_carton', 'net_weight', 'gross_weight', 'volume', 'price'}
IDENTIFIERS = {'po', 'customer_order', 'item', 'customer_item', 'customer_po', 'so'}
REQUIRED = {'po', 'item', 'quantity', 'cartons'}


def header_key(value):
    text = unicodedata.normalize('NFKC', str(value or '')).lower()
    return re.sub(r'[^a-z0-9\u4e00-\u9fff]', '', text)


def has_value(value):
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def heading_columns(sheet, row):
    return {FIELDS[header_key(cell.value)][0]: cell.column for cell in sheet[row]
            if header_key(cell.value) in FIELDS}


def table_columns(sheet, header_row):
    columns = []
    header_bottom = header_row
    used = set()
    for cell in sheet[header_row]:
        if not has_value(cell.value):
            continue
        key, label = FIELDS.get(header_key(cell.value), (f'column_{cell.column}', str(cell.value).strip()))
        if key in used:
            raise ValueError(f'{sheet.title} 第 {header_row} 行存在重复表头：{cell.value}')
        used.add(key)
        columns.append(CustomsColumn(key, label, cell.column))
        merge = next((area for area in sheet.merged_cells.ranges if cell.coordinate in area), None)
        if merge:
            header_bottom = max(header_bottom, merge.max_row)
            if key == 'description' and merge.max_col == cell.column + 1:
                columns.append(CustomsColumn('chinese_name', '中文品名', cell.column + 1))
            elif key == 'length' and merge.max_col == cell.column + 2:
                columns += [CustomsColumn('width', '宽(cm)', cell.column + 1),
                            CustomsColumn('height', '高(cm)', cell.column + 2)]
            elif merge.max_col > cell.column:
                # Preserve additional columns even for an unfamiliar merged heading.
                for col in range(cell.column + 1, merge.max_col + 1):
                    columns.append(CustomsColumn(f'column_{col}', f'{label}({get_column_letter(col)})', col))
    # Unit-only second header rows also occur when the vertical merges are absent.
    if header_bottom == header_row and header_row < sheet.max_row:
        below = [header_key(c.value) for c in sheet[header_row + 1] if has_value(c.value)]
        if below and all(value in {'kgs', 'kg', 'usd', 'cm', 'pcs', 'ctn'} for value in below):
            header_bottom += 1
    return sorted(columns, key=lambda col: col.column), header_bottom


def convert_cell(cell, raw, column, epoch, warnings):
    value = cell.value
    if cell.data_type == 'e':
        warnings.append(f'{column.label} 为 Excel 错误 {value}')
        return value
    if value is None:
        if raw.data_type == 'f':
            warnings.append(f'{column.label} 的公式没有缓存值，请用 Excel 重新计算并保存')
        return None
    if column.key == 'ship_date':
        if isinstance(value, (datetime, date)):
            return value.date() if isinstance(value, datetime) else value
        if isinstance(value, (int, float)):
            try:
                parsed = from_excel(value, epoch)
                if isinstance(parsed, datetime):
                    return parsed.date()
            except (ValueError, OverflowError):
                pass
            warnings.append('出货日期无法识别为日期')
        return value
    if column.key in NUMERIC:
        try:
            parsed = Decimal(str(value).replace(',', '').strip())
            if parsed.is_finite():
                return parsed
        except InvalidOperation:
            pass
        warnings.append(f'{column.label} 不是有效数字：{value}')
        return value
    if column.key in IDENTIFIERS and isinstance(value, (int, float)) and not isinstance(value, bool):
        if value == int(value):
            text = str(int(value))
            # Respect simple zero-padded number formats used for identifiers.
            if re.fullmatch(r'0+', cell.number_format):
                text = text.zfill(len(cell.number_format))
            return text
    return value.strip() if isinstance(value, str) else value


def read_customs_packing(path):
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise ValueError('海关装箱单请选择 .xlsx 文件')
    data = CustomsPackingData(path.resolve())
    raw_book = load_workbook(path, data_only=False)
    try:
        cached_book = load_workbook(path, data_only=True)
        try:
            for raw_sheet in raw_book.worksheets:
                sheet = cached_book[raw_sheet.title]
                header_row = next((r for r in range(1, min(raw_sheet.max_row, 50) + 1)
                                   if REQUIRED <= set(heading_columns(raw_sheet, r))), None)
                if header_row is None:
                    continue
                columns, bottom = table_columns(raw_sheet, header_row)
                for r in range(bottom + 1, raw_sheet.max_row + 1):
                    raw_cells = [raw_sheet.cell(r, col.column) for col in columns]
                    if not any(has_value(cell.value) for cell in raw_cells):
                        continue
                    texts = [header_key(cell.value) for cell in raw_cells if isinstance(cell.value, str)]
                    if any(text in {'total', 'totalusd', '合计', '总计'} for text in texts):
                        break
                    if any(text.startswith(('温馨提示', '出口方式', '出口国家', '车牌')) for text in texts):
                        break
                    if REQUIRED <= set(heading_columns(raw_sheet, r)):
                        continue  # Repeated page heading, not a shipment line.
                    warnings = []
                    values = {col.key: convert_cell(sheet.cell(r, col.column), raw_sheet.cell(r, col.column),
                                                   col, raw_book.epoch, warnings) for col in columns}
                    for key in REQUIRED:
                        if not has_value(values.get(key)):
                            name = next(col.label for col in columns if col.key == key)
                            warnings.append(f'{name}为空，已保留此行')
                    data.rows.append(CustomsPackingRow(sheet.title, r, columns, values,
                        {col.key: raw_sheet.cell(r, col.column).value for col in columns}, warnings))
                    data.warnings.extend(f'{sheet.title} 第 {r} 行：{warning}' for warning in warnings)
            if not data.rows:
                raise ValueError('未找到海关装箱单明细，请确认含有 PO#、Item#、Qty/PCS、no.of ctn 表头及货物数据')
            return data
        finally:
            cached_book.close()
    finally:
        raw_book.close()


def value_text(value):
    if value is None or value == '':
        return '未填写'
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, Decimal):
        # Keep full source precision in the model; remove floating-point noise only in the display.
        return format(value, '.6f').rstrip('0').rstrip('.')
    return str(value).strip()


def result_columns(data):
    """Use one shared header even when the workbook contains different sheet layouts."""
    columns = {}
    for row in data.rows:
        for column in row.columns:
            columns.setdefault(column.key, column)
    return list(columns.values())


def pipe_text(value):
    # Keep each record on one line and reserve ASCII | for the field separator.
    return ' '.join(value_text(value).split()).replace('|', '｜')


def header_message(columns):
    return ' | '.join(['序号', *(pipe_text(col.label) for col in columns)])


def row_message(index, row, columns=None):
    return ' | '.join([str(index), *(pipe_text(row.values.get(col.key)) for col in
                                    (columns if columns is not None else row.columns))])


def packing_summary(data):
    parts = [f'已读取 {len(data.rows)} 行货物明细']
    for key, label in (('cartons', '合计箱数'), ('quantity', '合计数量')):
        values = [row.values.get(key) for row in data.rows]
        if all(isinstance(value, Decimal) for value in values):
            parts.append(f'{label} {value_text(sum(values, Decimal(0)))}')
    parts.append(f'后续 ASN 将以这 {len(data.rows)} 行为基础逐行生成')
    return '；'.join(parts) + '。'
