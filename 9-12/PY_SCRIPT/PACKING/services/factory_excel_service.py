"""Write the factory notice using the supplied native Excel template."""
from copy import copy
from io import BytesIO
from pathlib import Path
import math
import unicodedata

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.views import Selection
from openpyxl.workbook.properties import CalcProperties

from models.packing import PackingDataError
from services.packing_lookup_service import number


FACTORY_REMINDER = ('温馨提示：\n'
                    '如外箱有箭头标识时，装货一定要箭头朝上，否则货到保税仓会有费用产生，谢谢。')
SHIPPING_FIELDS = ('订仓号：', '柜号：', '封条号：', '出口方式：',
                   '出口国家：', '出货期：', '运输：', '车牌：')


def add_factory_footer(sheet, total_row):
    """Anchor the reminder and blank shipping form below the actual totals."""
    start = total_row + 2
    end = start + len(SHIPPING_FIELDS) - 1
    for row in range(start, end + 1):
        sheet.row_dimensions[row].height = 26
    note_start = start + 1
    sheet.merge_cells(start_row=note_start, start_column=6,
                      end_row=note_start + 3, end_column=12)
    for cells in sheet.iter_rows(min_row=note_start, max_row=note_start + 3, min_col=6, max_col=12):
        for cell in cells:
            cell.fill = PatternFill('solid', fgColor='FFFF00')
    note = sheet.cell(note_start, 6, FACTORY_REMINDER)
    note.font = Font(name='宋体', size=14, bold=True, color='FF0000')
    note.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    border = Border(bottom=Side(style='thin', color='BFBFBF'))
    for r, title in enumerate(SHIPPING_FIELDS, start):
        sheet.merge_cells(start_row=r, start_column=16, end_row=r, end_column=17)
        sheet.merge_cells(start_row=r, start_column=18, end_row=r, end_column=25)
        label = sheet.cell(r, 16, title)
        label.font = Font(name='宋体', size=12)
        label.alignment = Alignment(horizontal='left', vertical='center')
        value = sheet.cell(r, 18)
        value.font = Font(name='宋体', size=12)
        value.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        for col in range(18, 26):
            sheet.cell(r, col).border = border
    return end


def write_factory_workbook(rows, template, directory):
    if not rows:
        raise PackingDataError('没有工厂箱单明细')
    total = sum(row.cartons for row in rows)
    if total <= 0 or total != total.to_integral_value():
        raise PackingDataError('总箱数必须为正整数')
    destination = Path(directory) / f'{int(total)}.xlsx'
    if destination.exists():
        raise FileExistsError(f'文件已存在，未覆盖：{destination}。请更换输出目录或移走旧文件后重试')
    book = load_workbook(template)
    try:
        sheet = book.worksheets[0]
        # Keep template title/header/column widths, replace every sample detail.
        detail_styles = [copy(sheet.cell(4, col)._style) for col in range(1, 26)]
        total_styles = [copy(sheet.cell(8, col)._style) for col in range(1, 26)]
        detail_height = sheet.row_dimensions[4].height or 45
        for merged in list(sheet.merged_cells.ranges):
            if merged.max_row >= 4:
                sheet.unmerge_cells(str(merged))
        sheet.delete_rows(4, max(1, sheet.max_row - 3))
        for row_no in list(sheet.row_dimensions):
            if row_no >= 4:
                del sheet.row_dimensions[row_no]
        # The reference includes two empty sheets; output has one notice only.
        for extra in list(book.worksheets[1:]):
            book.remove(extra)
        sheet.title = '工厂装箱单'
        # The reference hides price and has narrow numeric columns. Keep its
        # styling while exposing requested fields and avoiding Excel #### cells.
        widths = {'A': 13, 'B': 22, 'C': 17, 'E': 14, 'H': 10, 'I': 8, 'J': 10,
                  'K': 8, 'L': 8, 'M': 8, 'N': 9, 'O': 9, 'P': 13,
                  'Q': 13, 'R': 12, 'U': 25, 'V': 9, 'Y': 19}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
            sheet.column_dimensions[column].hidden = False
        for r, row in enumerate(rows, 4):
            item, values = row.item, row.product['values']
            data = [row.customer, row.order, item.customer_po, item.item_no.split('-')[0],
                    item.customer_item_no, item.description, str(values['产品名称']).strip(),
                    float(item.quantity), float(item.case_pack), f'=H{r}/I{r}',
                    *[float(number(values[k], k)) for k in ('长', '宽', '高', '整箱净重kg', '整箱毛重kg')],
                    f'=J{r}*N{r}', f'=J{r}*O{r}', f'=J{r}*K{r}*L{r}*M{r}/1000000',
                    f'=J{r}/{row.cartons_per_slip}' if row.cartons_per_slip else None,
                    item.description.split()[0], row.packaging, None, None, None, row.booking]
            for col, value in enumerate(data, 1):
                cell = sheet.cell(r, col, value)
                cell._style = copy(detail_styles[col - 1])
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                if col in (1, 2, 3, 4, 5, 6, 7, 20, 21, 25) and value is not None:
                    cell.data_type = 's'
                    cell.number_format = '@'
                elif col in (8, 9, 10, 19):
                    cell.number_format = '#,##0'
                elif col in (11, 12, 13, 14, 15, 16, 17, 18):
                    cell.number_format = '0.000'
            height = detail_height
            for col in (1, 2, 3, 5, 6, 7, 20, 21, 25):
                cell = sheet.cell(r, col)
                width = sheet.column_dimensions[cell.column_letter].width
                font_size = cell.font.sz or 11
                # Estimate wraps conservatively, accounting for wide Chinese text.
                capacity = max(1, (width - 1) * 11 / font_size)
                lines = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1
                                               for c in line) / capacity))
                            for line in str(cell.value or '').split('\n'))
                height = max(height, lines * font_size * 1.5 + 12)
            sheet.row_dimensions[r].height = height
        last = len(rows) + 3
        total_row = last + 1
        for col in range(1, 26):
            cell = sheet.cell(total_row, col)
            cell._style = copy(total_styles[col - 1])
            cell.alignment = Alignment(horizontal='center', vertical='center')
        sheet.cell(total_row, 2, 'TOTAL:')
        sheet.merge_cells(start_row=total_row, start_column=2, end_row=total_row, end_column=4)
        for col in ('H', 'J', 'P', 'Q', 'R', 'S'):
            cell = sheet[f'{col}{total_row}']
            cell.value = f'=SUM({col}4:{col}{last})'
            cell.number_format = '#,##0' if col in ('H', 'J', 'S') else '0.000'
        sheet.row_dimensions[total_row].height = 27
        footer_end = add_factory_footer(sheet, total_row)
        # Do not inherit the sample's B1 scroll origin or off-screen selection.
        sheet.sheet_view.topLeftCell = 'A1'
        sheet.freeze_panes = 'A4'
        sheet.sheet_view.selection = [Selection(pane='bottomLeft', activeCell='A4', sqref='A4')]
        book.active = 0
        sheet.print_area = f'A1:Y{footer_end}'
        sheet.print_title_rows = '1:3'
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')
        buffer = BytesIO()
        book.save(buffer)
        destination.parent.mkdir(parents=True, exist_ok=True)
        # Exclusive create also protects against another app creating the same name.
        with destination.open('xb') as stream:
            stream.write(buffer.getvalue())
        return destination
    finally:
        book.close()
