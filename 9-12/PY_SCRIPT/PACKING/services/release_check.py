"""Explicit --self-test checks for the packaged executable; never run during normal use."""
import io
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import traceback


def run_release_check(report_path):
    result = {'passed': False, 'checks': []}
    app = None
    try:
        from config import (PROJECT_DIR, RESOURCE_DIR, SETTINGS_FILE, PACKING_RULES_FILE, ASN_TEMPLATE,
                            ZHONGTONG_ASN_TEMPLATE, XINGHUI_ASN_TEMPLATE, PACKING_TEMPLATE, FACTORY_TEMPLATE)
        from services.settings_service import read_settings
        from ui.app import PackingApp
        app = PackingApp()
        app.withdraw()
        app.update_idletasks()
        assert app.title().casefold() == 'Love Anna, 加油!'.casefold()
        assert app.asn_tab.select_button.cget('text') == '选择三文件所在的目录'
        assert app.tabs.get() == '入仓预申报ASN生成器'
        assert set(app.asn_tab.warehouse_buttons) == {'yixing', 'zhongtong', 'xinghui'}
        assert 'customers' not in app.factory_tab.path_vars
        assert app.factory_tab.pdf_selection.selected_files() == ()
        assert SETTINGS_FILE.parent == PROJECT_DIR and SETTINGS_FILE.is_file()
        assert PACKING_RULES_FILE.is_file()
        result.update(app_dir=str(PROJECT_DIR), resource_dir=str(RESOURCE_DIR), window_title=app.title(),
                      font_size=read_settings(SETTINGS_FILE)['ui_font_size'])
        result['checks'].append('GUI, external JSON creation/loading, three warehouses, factory PDF selection without customer directory')

        import xlrd
        from openpyxl import load_workbook
        for template in (PACKING_TEMPLATE, FACTORY_TEMPLATE, ZHONGTONG_ASN_TEMPLATE, XINGHUI_ASN_TEMPLATE):
            book = load_workbook(template, read_only=True)
            assert book.sheetnames
            book.close()
        book = xlrd.open_workbook(str(ASN_TEMPLATE))
        assert '报关资料与ASN' in book.sheet_names()
        template_sheets = book.sheet_names()
        book.release_resources()
        from zipfile import ZipFile
        with ZipFile(XINGHUI_ASN_TEMPLATE) as archive:
            assert archive.read('xl/vbaProject.bin')
        result['checks'].append('All five embedded Excel templates and XLS/XLSX/XLSM readers, Xinghui VBA project')

        from pypdf import PdfReader, PdfWriter
        pdf = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(pdf)
        pdf.seek(0)
        assert len(PdfReader(pdf).pages) == 1
        result['checks'].append('PDF reader/writer')

        from services.asn_generation_service import write_asn_workbook
        row = [None] * 47
        row[0], row[1], row[5], row[26], row[28] = 1, 'RELEASE-CHECK', '00123', 8, 7.5
        plan = {'template': str(ASN_TEMPLATE), 'sheet': '报关资料与ASN', 'scope': 'all',
                'contract': 'RELEASE-CHECK', 'headers': {'C8': 'Release test customer'}, 'rows': [row],
                'formulas': [{'cell': 'AD26', 'value': '=ROUND(AA26*AC26,2)'}]}
        with TemporaryDirectory(prefix='love-anna-check-') as temp:
            output = write_asn_workbook(plan, temp)
            book = xlrd.open_workbook(str(output))
            assert book.sheet_names() == template_sheets
            sheet = book.sheet_by_name('报关资料与ASN')
            assert sheet.cell_value(7, 2) == 'Release test customer'
            assert sheet.cell_value(25, 5) == '00123'
            assert sheet.cell_value(25, 29) == 60
            book.release_resources()
        result['checks'].append('Embedded PowerShell writer and native Excel XLS generation')

        for warehouse, template, sheet_name, amount, quantity, price, columns in (
            ('zhongtong', ZHONGTONG_ASN_TEMPLATE, 'ASN', 'P', 'M', 'O', 30),
            ('xinghui', XINGHUI_ASN_TEMPLATE, '报关资料与ASN', 'U', 'R', 'T', 47),
        ):
            from openpyxl.utils import column_index_from_string
            row = [None] * columns
            row[0], row[1] = 1, 'RELEASE-CHECK'
            row[column_index_from_string(quantity) - 1] = 8
            row[column_index_from_string(price) - 1] = 7.5
            plan = {'warehouse': warehouse, 'template': str(template), 'sheet': sheet_name,
                    'scope': 'all', 'contract': 'RELEASE-CHECK', 'headers': {}, 'rows': [row],
                    'formulas': [{'cell': f'{amount}26', 'value': f'=ROUND({quantity}26*{price}26,2)'}]}
            with TemporaryDirectory(prefix='love-anna-check-') as temp:
                output = write_asn_workbook(plan, temp)
                original = load_workbook(template, read_only=True)
                saved = load_workbook(output, read_only=True, data_only=True)
                try:
                    assert original.sheetnames == saved.sheetnames
                    assert saved[sheet_name][f'{amount}26'].value == 60
                finally:
                    original.close()
                    saved.close()
                if warehouse == 'xinghui':
                    with ZipFile(template) as source, ZipFile(output) as exported:
                        assert source.read('xl/vbaProject.bin') == exported.read('xl/vbaProject.bin')
            result['checks'].append(f'Native {warehouse} generation, all sheets retained' +
                                    (', original VBA bytes retained' if warehouse == 'xinghui' else ''))
        result['passed'] = True
    except Exception:
        result['error'] = traceback.format_exc()
    finally:
        if app is not None:
            for timer in app.tk.call('after', 'info'):
                app.tk.call('after', 'cancel', timer)
            app.destroy()
        Path(report_path).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    return 0 if result['passed'] else 1
