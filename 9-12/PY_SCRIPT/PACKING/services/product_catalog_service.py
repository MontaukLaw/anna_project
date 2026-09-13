"""只读预加载产品资料，跳过图片资源，将产品行缓存到内存。"""
from pathlib import Path
from queue import Queue
import re

from openpyxl import load_workbook

from models.product_catalog import ProductCatalog
from services.settings_service import save_product_path


def read_product_catalog(path: Path) -> ProductCatalog:
    if path.suffix.lower() != ".xlsx":
        raise ValueError("请选择 .xlsx 格式的 JP 产品资料表")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = {}
    items = {}
    try:
        for sheet in workbook:
            rows = list(sheet.iter_rows(values_only=True))
            sheets[sheet.title] = rows
            for row_number, row in enumerate(rows[:30]):
                index = next((i for i, value in enumerate(row)
                              if re.sub(r"[^a-z0-9]", "", str(value).casefold()) in {"itemno", "itemnumber"}), None)
                if index is None:
                    continue
                headers = [str(value).strip() if value is not None else f"列{i+1}" for i, value in enumerate(row)]
                for source_row, values in enumerate(rows[row_number + 1:], start=row_number + 2):
                    value = values[index]
                    if value is None or not str(value).strip():
                        continue
                    item = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()
                    items.setdefault(item, []).append({
                        "sheet": sheet.title, "row": source_row,
                        "values": dict(zip(headers, values)),
                    })
                break
    finally:
        workbook.close()
    if not items:
        raise ValueError("资料表中未找到带有 ITEM NO 的产品数据，请选择正确的 JP 产品资料表。")
    return ProductCatalog(path, sheets, items)


def load_product_catalog_job(path: Path, config_path: Path, remember: bool, events: Queue):
    try:
        catalog = read_product_catalog(path)
        events.put(("product_loaded", catalog))
        if remember:
            try:
                save_product_path(config_path, path)
            except Exception as exc:
                events.put(("product_config_error", str(exc)))
            else:
                events.put(("product_remembered", str(config_path)))
    except Exception as exc:
        events.put(("product_error", str(exc)))
    finally:
        events.put(("product_done", None))
