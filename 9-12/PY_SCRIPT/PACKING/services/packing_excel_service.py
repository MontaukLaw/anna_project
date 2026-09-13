"""按范例样式生成装箱单，体积、重量与合计使用 Excel 公式。"""
from copy import copy
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter
from services.slip_summary_service import append_slip_summaries


def build_packing_excel(template, order, customer, rows):
    workbook = load_workbook(template)
    sheet = workbook.active
    original_end = sheet.max_row
    snapshots = {}
    for row in sheet.iter_rows(min_row=20):
        for cell in row:
            if not isinstance(cell, MergedCell):
                snapshots[(cell.row, cell.column)] = (cell.value, copy(cell._style), copy(cell.alignment), cell.number_format)
    heights = {r: sheet.row_dimensions[r].height for r in range(20, original_end + 1)}
    merges = [copy(region) for region in sheet.merged_cells.ranges if region.min_row >= 20]
    for region in merges:
        sheet.unmerge_cells(str(region))
    sheet.delete_rows(20, original_end - 19)
    for key in list(sheet.row_dimensions):
        if key >= 20:
            del sheet.row_dimensions[key]
    count = len(rows)
    total_row = 20 + count
    slip_notes = list(dict.fromkeys(r.slip_note for r in rows if r.slip_note))
    show_totals = not slip_notes
    shift = count - 4 + len(slip_notes) - int(not show_totals)

    def copy_row(source, target, values=True):
        for col in range(1, 28):
            info = snapshots.get((source, col))
            if info:
                cell = sheet.cell(target, col)
                cell.value = info[0] if values else None
                cell._style = copy(info[1])
                cell.alignment = copy(info[2])
                cell.number_format = info[3]
        sheet.row_dimensions[target].height = heights.get(source)

    for old in range(25, original_end + 1):
        copy_row(old, old + shift)
    for region in merges:
        region.shift(row_shift=shift)
        sheet.merge_cells(str(region))
    totals = {5: Decimal(0), 6: Decimal(0), 19: Decimal(0), 21: Decimal(0), 23: Decimal(0)}
    for row_index, row in enumerate(rows, 20):
        copy_row(20, row_index, values=False)
        quantity = row.display_quantity if row.display_quantity is not None else row.item.quantity
        case_pack = row.display_case_pack if row.display_case_pack is not None else row.item.case_pack
        ct = quantity / case_pack
        box = row.length * row.width * row.height / Decimal(1000000)
        volume = (box * ct).quantize(Decimal('0.001'), rounding=ROUND_HALF_UP)
        values = [order, row.item.customer_po, row.item.item_no, case_pack, quantity,
                  ct, row.net, 'KGS', row.gross, 'KGS', row.length, 'x', row.width, 'x', row.height, 'CM',
                  box, 'CBM', volume, 'CBM', row.net * ct, 'KGS', row.gross * ct, 'KGS', row.english_name, 'N']
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row_index, col, float(value) if isinstance(value, Decimal) else value)
            if isinstance(value, str):
                cell.data_type = 's'
            if col in (1, 2, 3):
                cell.number_format = '@'
        sheet.cell(row_index, 3).comment = Comment(f"DATE CODE: {row.date_code}\n资料来源: {row.source}", "Packing System")
        for col, formula in {
            17: f'=K{row_index}*M{row_index}*O{row_index}/1000000',
            19: f'=Q{row_index}*F{row_index}',
            21: f'=G{row_index}*F{row_index}',
            23: f'=I{row_index}*F{row_index}',
        }.items():
            cell = sheet.cell(row_index, col, formula)
            cell.number_format = '0.000' if col in (17, 19) else '0.00'
        if row.cargo_label:
            sheet.cell(row_index, 27, row.cargo_label)
            sheet.cell(row_index, 27).alignment = Alignment(wrap_text=True, vertical='center')
            sheet.row_dimensions[row_index].height = max(sheet.row_dimensions[row_index].height or 20, 32)
        for col in totals:
            if row.cargo_label != 'cargo':
                totals[col] += row.item.quantity if col == 5 else values[col - 1]
    if show_totals:
        copy_row(24, total_row, values=False)
        for col, value in totals.items():
            letter = get_column_letter(col)
            cell = sheet.cell(total_row, col, f'=SUM({letter}20:{letter}{total_row-1})')
            cell.number_format = '0.000' if col == 19 else ('0.00' if col in (21,23) else '0')
        for col, value in {20: 'CBM', 22: 'KGS', 24: 'KGS'}.items():
            sheet.cell(total_row, col, value)
    for offset, note in enumerate(slip_notes, 1):
        index = total_row + offset - int(not show_totals)
        sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=26)
        cell = sheet.cell(index, 1, note)
        cell.fill = PatternFill('solid', fgColor='FFF2CC')
        cell.font = Font(bold=True, color='C00000', size=12)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[index].height = 30
    sheet['U6'] = datetime.now().date()
    sheet['T8'] = order
    sheet['T10'] = ' / '.join(dict.fromkeys(r.item.customer_po for r in rows))
    sheet['T12'] = customer
    dates = list(dict.fromkeys(r.delivery_date for r in rows))
    if any(value is None for value in dates):
        raise ValueError('缺少订单 PDF Cargo Ready Date')
    sheet['C12'] = dates[0] if len(dates) == 1 else ' / '.join(value.isoformat() for value in dates)
    sheet['C12'].number_format = 'yyyy-mm-dd'
    codes = list(dict.fromkeys(r.date_code for r in rows))
    sheet['T14'] = codes[0] if len(codes) == 1 else '\n'.join(f"{r.item.item_no.split('-')[0]}: {r.date_code}" for r in rows)
    if len(codes) > 1:
        sheet.merge_cells('T14:Z14')
        alignment = copy(sheet['T14'].alignment)
        alignment.wrap_text = True
        sheet['T14'].alignment = alignment
        sheet.row_dimensions[14].height = max(30, 15 * len(rows))
    has_slip = bool(slip_notes)
    if has_slip:
        sheet.cell(19, 27, 'REMARKS')
        sheet.column_dimensions['AA'].width = 32
    end_row = original_end + shift
    if has_slip:
        end_row = append_slip_summaries(sheet, rows, end_row + 2)
    sheet.print_area = f'A1:{"AA" if has_slip else "Z"}{end_row}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet._images = []
    workbook.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
