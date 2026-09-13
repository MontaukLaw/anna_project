"""在装箱单末尾写入每个滑托 ITEM 的重量拆分和包装统计。"""
from decimal import Decimal
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from models.packing import PackingDataError


def append_slip_summaries(sheet, rows, start_row):
    index = start_row
    for position, pallet in enumerate(rows):
        if pallet.cargo_label != 'cargo together with slip sheet':
            continue
        cargo = rows[position - 1] if position else None
        if cargo is None or cargo.cargo_label != 'cargo' or cargo.item.item_no != pallet.item.item_no:
            raise PackingDataError('SLIP SHEET 明细缺少对应的 cargo 行')
        per_slip = pallet.display_case_pack
        product_weight = cargo.gross * per_slip
        other_weight = pallet.gross - product_weight
        if other_weight < 0:
            raise PackingDataError(f'{pallet.item.item_no}：整托毛重小于单箱毛重乘以每托箱数，请核对资料表')
        details = [
            ('ITEM #', pallet.item.item_no, ''),
            ('Product weight + Carton Box Weight :', product_weight, 'KG'),
            ('Total others packing material weight as 彩盒，说明书，卡纸类:', other_weight, 'KG'),
            ('GRAND TOTAL CARTON GROSS WEIGHT PER CARTON :', pallet.gross, 'KG'),
            ('訂單該 ITEM 的 slip sheet 數量', pallet.display_quantity / per_slip, '托'),
            ('每 slip sheet 的 CTN 數量', per_slip, '箱'),
            ('每 slip sheet 的總體積', pallet.length * pallet.width * pallet.height / Decimal(1000000), 'CBM'),
            ('每 slip sheet 的淨重', pallet.net, 'KG'),
            ('每 slip sheet 的毛重', pallet.gross, 'KG'),
        ]
        for offset, (title, value, unit) in enumerate(details):
            row = index + offset
            ranges = [(1, 6), (7, 12)] if offset == 0 else [(1, 6), (7, 9), (10, 12)]
            for first, last in ranges:
                sheet.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
            sheet.cell(row, 1, title)
            cell = sheet.cell(row, 7, float(value) if isinstance(value, Decimal) else value)
            if isinstance(value, str):
                cell.data_type = 's'
            cell.number_format = ('0' if unit in ('托', '箱') else ('0.000000' if unit == 'CBM' else '0.00')) if isinstance(value, Decimal) else '@'
            if offset:
                sheet.cell(row, 10, unit)
            for col in range(1, 13):
                target = sheet.cell(row, col)
                color = 'E8EEF5' if offset == 0 else ('F4F7FA' if offset % 2 == 0 else 'FFFFFF')
                if offset == 3:
                    color = 'FFF2CC'
                target.fill = PatternFill('solid', fgColor=color)
                target.border = Border(bottom=Side(style='thin', color='DCE3EB'))
            for col in (1, 7) if offset == 0 else (1, 7, 10):
                target = sheet.cell(row, col)
                target.font = Font(name='Arial', size=11, bold=offset in (0, 3), color='243746')
                target.alignment = Alignment(vertical='center', wrap_text=True,
                    horizontal='right' if col == 7 and offset else 'left', indent=1)
            sheet.row_dimensions[row].height = 42 if offset in (2, 3) else 28
        index += len(details) + 1
    return index - 1
