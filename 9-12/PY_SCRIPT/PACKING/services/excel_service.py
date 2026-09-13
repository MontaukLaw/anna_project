from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from models.table import TableData


def write_table(table: TableData, destination: Path) -> Path:
    """将识别结果写为文本单元格，保留前导零并避免文本被当作公式。"""
    table.validate()
    if destination.suffix.lower() != ".xlsx":
        raise ValueError("Excel 文件扩展名必须是 .xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "箱单明细"
    for row_index, row in enumerate([table.columns, *table.rows], start=1):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, column_index, str(value))
            cell.data_type = "s"
            cell.number_format = "@"
            if row_index == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="334155")
    for index, title in enumerate(table.columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(40, max(16, len(title) + 4))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    try:
        workbook.save(destination)
    finally:
        workbook.close()
    return destination
