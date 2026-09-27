"""Read ASN source documents without modifying them; retain cell references in the audit log."""
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from services.asn_quantity_service import quantities_with_units
from services.hs_units_service import unit_name


KINDS = ('装箱单', '香港合同', '形式发票')
FILE_PATTERN = re.compile(r'^(装箱单|香港合同|销售合同|形式发票)\s*(.+)\.(xls|xlsx)$', re.I)
MONEY = Decimal('0.01')


def normalized(value):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', '' if value is None else str(value)))


def display(value):
    if value is None or value == '':
        return '未读取到'
    if isinstance(value, Decimal):
        return format(value, 'f').rstrip('0').rstrip('.') if '.' in format(value, 'f') else str(value)
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    return str(value).strip()


def number(value):
    if value is None or isinstance(value, (bool, date)):
        return None
    text = normalized(value).replace(',', '')
    text = re.sub(r'^(USD|US\$|\$)', '', text, flags=re.I)
    try:
        result = Decimal(text)
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


@dataclass
class Sheet:
    name: str
    rows: list[list]

    def ref(self, row, column):
        return f'{self.name}!{get_column_letter(column + 1)}{row + 1}'


@dataclass
class Item:
    name: str
    quantity: Decimal | None
    price: Decimal | None
    amount: Decimal | None
    source: str
    code: str = ''
    cartons: Decimal | None = None
    per_carton: Decimal | None = None
    country: str = ''
    specification: str = ''
    unit: str = ''
    net_weight: Decimal | None = None
    gross_weight: Decimal | None = None
    quantity_units: dict[str, Decimal] = field(default_factory=dict)
    quantity_text: str = ''
    quantity_unit: str = ''


@dataclass
class Document:
    kind: str
    path: Path
    contract: str = ''
    signed_date: str = ''
    items: list[Item] = field(default_factory=list)
    total: Decimal | None = None
    currency: str = ''

    @property
    def quantity(self):
        if not self.items or any(item.quantity is None for item in self.items):
            return None
        return sum((item.quantity for item in self.items), Decimal(0))


@dataclass
class AuditReport:
    entries: list[tuple[str, str]] = field(default_factory=list)
    documents: dict[str, dict[str, Document]] = field(default_factory=dict)
    checks: int = 0
    failures: int = 0
    issues: list[tuple[str, str]] = field(default_factory=list)
    contract_failures: dict[str, int] = field(default_factory=dict)

    @property
    def passed(self):
        return self.checks > 0 and self.failures == 0


class Auditor:
    def __init__(self, emit=None):
        self.report = AuditReport()
        self.emit = emit

    def log(self, text, level='INFO'):
        self.report.entries.append((text, level))
        if self.emit:
            self.emit(text, level)

    def check(self, condition, title, details):
        self.report.checks += 1
        if not condition:
            self.report.failures += 1
            self.report.issues.append((title, details))
        self.log(f'{"一致 / 通过" if condition else "不一致 / 未通过"}：{title}；{details}',
                 'SUCCESS' if condition else 'ERROR')


def read_sheets(path):
    if path.suffix.lower() == '.xls':
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError('读取 .xls 需要 xlrd，请安装 requirement.txt 中的依赖') from exc
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            sheets = []
            for sheet in book.sheets():
                rows = []
                for r in range(sheet.nrows):
                    values = []
                    for cell in sheet.row(r):
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, book.datemode)
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            value = xlrd.error_text_from_code.get(value, '#ERROR')
                        values.append(value)
                    rows.append(values)
                sheets.append(Sheet(sheet.name, rows))
            return sheets
        finally:
            book.release_resources()
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return [Sheet(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
                for sheet in book.worksheets]
    finally:
        book.close()


def header_field(value):
    # Use the first (Chinese) line; contract headings include numbered English translations.
    text = re.sub(r'^\d+[.、]', '', normalized(str(value or '').split('\n')[0]))
    full = normalized(value)
    for field_name, aliases in (
        ('name', ('货物名称', '货名', '品名及规格', '品名')),
        ('code', ('商品编码', '海关编码')),
        ('quantity', ('总数量', '数量')),
        ('cartons', ('总箱数', '箱数')),
        ('per_carton', ('每箱数量',)),
        ('country', ('消费国',)),
        ('net_weight', ('总净重', '净重')),
        ('gross_weight', ('总毛重', '毛重')),
        ('unit', ('单位',)),
        ('price', ('单价',)),
        ('amount', ('总价', '金额')),
        ('specification', ('规格', '申报要素')),
    ):
        if any(text == alias or text.startswith(alias + '(') or full == alias for alias in aliases):
            return field_name
    return None


def find_table(sheets, kind):
    candidates = []
    for sheet in sheets:
        for r, row in enumerate(sheet.rows):
            mapping = {c: header_field(value) for c, value in enumerate(row) if header_field(value)}
            if 'name' not in mapping.values():
                continue
            if kind == '装箱单' and 'code' not in mapping.values():
                continue
            if kind != '装箱单' and not {'quantity', 'price', 'amount'} <= set(mapping.values()):
                continue
            end = r
            if kind == '装箱单':
                for rr in range(r + 1, min(r + 4, len(sheet.rows))):
                    fields = {c: header_field(v) for c, v in enumerate(sheet.rows[rr]) if header_field(v)}
                    if not fields:
                        break
                    mapping.update(fields)  # Lower header wins: I=单价, N=总数量 in the supplied sample.
                    end = rr
            columns = {}
            for c, key in mapping.items():
                if key in columns:
                    raise ValueError(f'{sheet.ref(r, c)}：{key} 表头重复，无法确定读取列')
                columns[key] = c
            candidates.append((sheet, r, end + 1, columns))
    if len(candidates) != 1:
        raise ValueError(f'找到 {len(candidates)} 个可识别的{kind}明细表，需要唯一的明细表（请核对表头）')
    return candidates[0]


def labeled_values(sheet, pattern, stop=None):
    result = []
    for r, row in enumerate(sheet.rows[:stop]):
        for c, value in enumerate(row):
            match = pattern.match(normalized(value))
            if not match:
                continue
            tail = normalized(value)[match.end():].lstrip(':')
            if tail and tail.upper() not in ('NO.', 'NO', 'DATE', 'USD'):
                result.append((tail, sheet.ref(r, c)))
            else:
                for cc in range(c + 1, len(row)):
                    if row[cc] is not None and str(row[cc]).strip():
                        result.append((row[cc], sheet.ref(r, cc)))
                        break
    return result


CONTRACT_LABEL = re.compile(r'^(?:协议合同号|合同协议号|合同号码|合同编号|合同号)(?::)?(?:NO\.?)?', re.I)
DATE_LABEL = re.compile(r'^(?:合同签订日期|合同签约日期|日期)(?::)?(?:DATE)?', re.I)
TOTAL_LABEL = re.compile(r'^(?:TOTAL:USD|TOTALAMOUNT|TOTALVALUE|(?:\d+\.)?总值:?USD|合计(?:\(USD\))?)$', re.I)


def read_document(path, kind, expected, audit):
    sheets = read_sheets(path)
    # Check identity immediately, before parsing any item values.
    identities = [(value, ref) for sheet in sheets
                  for value, ref in labeled_values(sheet, CONTRACT_LABEL)]
    audit.check(bool(identities) and all(normalized(value) == expected for value, _ in identities),
                f'{path.name} 合同号核对',
                f'文件名合同号={expected}；表内=' + ('、'.join(f'{display(v)} ({ref})' for v, ref in identities) or '未找到合同号'))
    sheet, header_row, start, columns = find_table(sheets, kind)
    doc = Document(kind, path, normalized(identities[0][0]) if identities else '')
    audit.log(f'读取 {path.name} / {sheet.name}，明细从第 {start + 1} 行开始')
    if kind == '香港合同':
        dates = labeled_values(sheet, DATE_LABEL, header_row)
        parsed = []
        for value, ref in dates:
            text = display(value)
            try:
                normalized_date = re.sub(r'[年/月.]', '-', text).replace('日', '')
                parsed.append(datetime.strptime(normalized_date, '%Y-%m-%d').date().isoformat())
            except ValueError:
                pass
            audit.log(f'合同签订日期：{text} ({ref})')
        audit.check(bool(parsed) and len(parsed) == len(dates) and len(set(parsed)) == 1,
                    f'{path.name} 合同签订日期', '、'.join(parsed) or '缺失或日期格式无法识别')
        doc.signed_date = parsed[0] if parsed else ''

    # A declaration next to 币别/币种 takes precedence over USD column headings.
    currencies = labeled_values(sheet, re.compile(r'^(?:币别|币种):?'), header_row)
    usd_aliases = {'USD', '美元', '美金', 'US$'}
    header_text = ''.join(normalized(v).upper() for row in sheet.rows[header_row:start] for v in row)
    packing_usd = any(normalized(v).upper() == 'TOTAL:USD' for row in sheet.rows for v in row)
    doc.currency = 'USD' if ((currencies and all(normalized(v).upper() in usd_aliases for v, _ in currencies))
                             or (not currencies and ('USD' in header_text or packing_usd))) else ''
    audit.check(doc.currency == 'USD', f'{path.name} 币种',
                '、'.join(display(v) for v, _ in currencies) if currencies else (doc.currency or '未识别到 USD'))

    total_rows = []
    for r in range(start, len(sheet.rows)):
        row = sheet.rows[r]
        if any(TOTAL_LABEL.match(normalized(v)) for v in row):
            total_rows.append(r)
    end = min(total_rows) if total_rows else len(sheet.rows)
    if not total_rows:
        audit.check(False, f'{path.name} 合计行', '未找到合计，不能确认明细边界')
    required = ['name', 'quantity', 'price', 'amount']
    if kind == '装箱单':
        required += ['code', 'cartons', 'country']
    audit.check(all(key in columns for key in required), f'{path.name} 必需表头',
                '完整' if all(key in columns for key in required) else '缺少：' + '、'.join(k for k in required if k not in columns))
    for r in range(start, end):
        row = sheet.rows[r]
        values = {key: row[c] if c < len(row) else None for key, c in columns.items()}
        if not any(value is not None and str(value).strip() for value in values.values()):
            continue
        # Don't silently discard partially populated item rows.
        item = Item(str(values.get('name') or '').strip(), number(values.get('quantity')),
                    number(values.get('price')), number(values.get('amount')), sheet.ref(r, columns['name']))
        item.quantity_text = str(values.get('quantity') or '')
        if item.quantity is None:
            item.quantity_units = quantities_with_units(values.get('quantity'))
            candidates = [(unit, qty) for unit, qty in item.quantity_units.items()
                          if item.price is not None and money_equal(qty * item.price, item.amount)]
            if len(candidates) == 1:
                item.quantity_unit, item.quantity = candidates[0]
                audit.log(f'{path.name} 第 {r + 1} 行计价数量：{display(item.quantity)}{item.quantity_unit}；'
                          f'按单价×数量与金额唯一匹配，保留全部单位数量：{item.quantity_units}')
            elif item.quantity_units:
                audit.check(False, f'{path.name} 第 {r + 1} 行计价数量',
                            f'总数量={item.quantity_text}，无法按单价与金额唯一确定计价数量，请核实')
        for key in ('code', 'country', 'unit', 'specification'):
            raw = values.get(key)
            setattr(item, key, str(int(raw)) if isinstance(raw, float) and raw.is_integer() else str(raw or '').strip())
        # The contract template's merged 品名及规格 header spans name and declaration columns.
        name_col = columns['name']
        if kind == '香港合同' and not item.specification and columns.get('quantity', 0) > name_col + 1:
            item.specification = ' / '.join(str(v).strip() for v in row[name_col + 1:columns['quantity']] if v is not None and str(v).strip())
        for key in ('cartons', 'per_carton', 'net_weight', 'gross_weight'):
            setattr(item, key, number(values.get(key)))
        details = [f'{key}={display(value)} ({sheet.ref(r, columns[key])})' for key, value in values.items()]
        names = {'name': '品名', 'quantity': '数量', 'price': '单价', 'amount': '金额', 'code': '商品编码',
                 'cartons': '箱数', 'per_carton': '每箱数量', 'country': '消费国', 'unit': '单位',
                 'specification': '规格', 'net_weight': '净重', 'gross_weight': '毛重'}
        audit.log(f'{path.name} 第 {r + 1} 行：' + '；'.join(
            detail.replace(key + '=', names[key] + '=', 1) for key, detail in zip(values, details)))
        if item.specification and 'specification' not in columns:
            audit.log(f'品名及规格补充：{item.specification} ({sheet.ref(r, name_col + 1)})')
        valid = bool(item.name) and all(v is not None and v >= 0 for v in (item.quantity, item.price, item.amount))
        audit.check(valid, f'{path.name} 第 {r + 1} 行明细完整性', '品名、数量、单价、金额须有效，缺失值不会当作 0')
        if valid:
            audit.check(money_equal(item.quantity * item.price, item.amount), f'{path.name} 第 {r + 1} 行 数量 × 单价',
                        f'{display(item.quantity)} × {display(item.price)} = {display(item.quantity * item.price)}；金额={display(item.amount)}')
        if kind == '装箱单':
            packing_valid = (item.code.isdigit() and bool(item.country) and not item.country.startswith('#')
                             and item.cartons is not None and item.cartons > 0
                             and item.cartons == item.cartons.to_integral_value())
            audit.check(packing_valid, f'{path.name} 第 {r + 1} 行装箱资料',
                        f'商品编码={item.code or "缺失"}（须为数字）；消费国={item.country or "缺失"}；箱数={display(item.cartons)}')
            if item.per_carton is None and 'per_carton' not in columns and item.quantity is not None and item.cartons and item.cartons > 0:
                item.per_carton = item.quantity / item.cartons
                audit.log(f'每箱数量（推算）：{display(item.quantity)} ÷ {display(item.cartons)} = {display(item.per_carton)}；原表未提供独立的每箱数量列，混装时需人工确认。')
            audit.check(item.per_carton is not None and item.per_carton > 0
                        and item.per_carton == item.per_carton.to_integral_value()
                        and item.cartons is not None and item.quantity is not None
                        and item.per_carton * item.cartons == item.quantity,
                        f'{path.name} 第 {r + 1} 行 每箱数量 × 箱数',
                        f'{display(item.per_carton)} × {display(item.cartons)}；总数量={display(item.quantity)}')
        doc.items.append(item)
    audit.check(bool(doc.items), f'{path.name} 商品明细', f'读取到 {len(doc.items)} 行')

    # Packing's 合计 row has quantities/weights; only TOTAL:USD is a monetary total.
    total_pattern = re.compile(r'^TOTAL:USD$', re.I) if kind == '装箱单' else TOTAL_LABEL
    totals = labeled_values(sheet, total_pattern)
    amount_values = [(number(value), ref) for value, ref in totals]
    doc.total = amount_values[0][0] if amount_values else None
    audit.log(f'{path.name} 合计金额：' + ('；'.join(f'{display(value)} ({ref})' for value, ref in amount_values) or '未读取到'))
    valid_totals = bool(amount_values) and all(value is not None for value, _ in amount_values)
    audit.check(valid_totals and all(money_equal(value, doc.total) for value, _ in amount_values),
                f'{path.name} 合计金额有效且一致', 'USD ' + display(doc.total))
    summed = sum((item.amount for item in doc.items), Decimal(0)) if doc.items and all(item.amount is not None for item in doc.items) else None
    audit.check(money_equal(summed, doc.total), f'{path.name} 明细金额合计',
                f'明细合计={display(summed)}；表内合计={display(doc.total)}')
    if kind == '装箱单':
        for r in total_rows:
            row = sheet.rows[r]
            if not any(normalized(v) == '合计' for v in row):
                continue
            for key, title in (('quantity', '总数量'), ('cartons', '总箱数')):
                col = columns.get(key)
                stated = number(row[col]) if col is not None and col < len(row) else None
                parts = [getattr(item, key) for item in doc.items]
                calculated = sum(parts, Decimal(0)) if parts and all(v is not None for v in parts) else None
                audit.check(stated is not None and stated == calculated, f'{path.name} {title}合计',
                            f'明细={display(calculated)}；合计行={display(stated)} ({sheet.ref(r, col) if col is not None else "缺少列"})')
    audit.log(f'{path.name} 汇总：总数量={display(doc.quantity)}；USD {display(doc.total)}')
    return doc


def money_equal(left, right):
    return left is not None and right is not None and left.quantize(MONEY, rounding=ROUND_HALF_UP) == right.quantize(MONEY, rounding=ROUND_HALF_UP)


def compare_documents(documents, contract, audit):
    if len(documents) != 3:
        audit.check(False, f'合同 {contract} 三单交叉核对', '文件缺失、重复或读取失败，无法完成核对')
        return
    docs = [documents[kind] for kind in KINDS]
    audit.check(all(doc.quantity is not None and doc.quantity == docs[0].quantity for doc in docs),
                f'合同 {contract} 三单总数量', '；'.join(f'{doc.kind}={display(doc.quantity)}' for doc in docs))
    audit.check(all(money_equal(doc.total, docs[0].total) for doc in docs),
                f'合同 {contract} 三单总金额（装箱单 TOTAL:USD / 销售合同 / 形式发票）',
                '；'.join(f'{doc.kind}=USD {display(doc.total)}' for doc in docs))
    groups = []
    for doc in docs:
        grouped = defaultdict(list)
        for item in doc.items:
            grouped[normalized(item.name)].append(item)
        groups.append(grouped)
    audit.check(bool(groups[0]) and set(groups[0]) == set(groups[1]) == set(groups[2]),
                f'合同 {contract} 三单品名', '；'.join(f'{doc.kind}={"、".join(group)}' for doc, group in zip(docs, groups)))
    for name in sorted(set().union(*(set(group) for group in groups))):
        if not all(name in group for group in groups):
            audit.check(False, f'商品 {name} 三单明细', '缺少该品名：' + '、'.join(doc.kind for doc, group in zip(docs, groups) if name not in group))
            continue
        # Aggregate split rows, but keep each price separate so swapped prices can't pass by total alone.
        signatures = []
        for doc, group in zip(docs, groups):
            quantities = defaultdict(Decimal)
            amounts = defaultdict(Decimal)
            valid = True
            for item in group[name]:
                if any(value is None for value in (item.price, item.quantity, item.amount)):
                    valid = False
                    break
                quantities[item.price] += item.quantity
                amounts[item.price] += item.amount
            signatures.append((dict(quantities), {price: amount.quantize(MONEY, rounding=ROUND_HALF_UP) for price, amount in amounts.items()}) if valid else None)
            audit.log(f'{doc.kind} / {name}：' + '；'.join(f'单价 {display(price)}，数量 {display(qty)}，金额 {display(amounts[price])}' for price, qty in quantities.items()))
        audit.check(all(sig is not None for sig in signatures) and signatures[0] == signatures[1] == signatures[2],
                    f'商品 {name} 数量、单价、金额', '按品名及单价汇总比较，详见上方各单据数值')
        for item in groups[0][name]:
            if item.quantity_units:
                units = {unit_name(entry.unit) for entry in groups[2][name] if entry.price == item.price}
                audit.check(units == {item.quantity_unit}, f'商品 {name} 多单位计价数量',
                            f'装箱单计价单位={item.quantity_unit}；发票单位={"、".join(sorted(units))}；'
                            f'总数量原文={item.quantity_text}')


def audit_directory(directory, emit=None):
    audit = Auditor(emit)
    directory = Path(directory)
    audit.log(f'开始核对目录：{directory}（仅检查当前目录，按文件名合同号分组）')
    try:
        files = sorted((p for p in directory.iterdir() if p.is_file()), key=lambda p: p.name.casefold())
    except OSError as exc:
        audit.check(False, '目录读取', str(exc))
        return audit.report
    groups = defaultdict(lambda: defaultdict(list))
    for path in files:
        if path.name.startswith('~$'):
            continue
        match = FILE_PATTERN.fullmatch(path.name)
        if match:
            kind, contract, _ = match.groups()
            contract = normalized(contract)
            if contract:
                groups[contract]['香港合同' if kind == '销售合同' else kind].append(path)
    if not groups:
        audit.check(False, '所需文件', '未找到 形式发票<合同号>.xls、装箱单<合同号>.xls、香港合同<合同号>.xls（也支持 .xlsx）')
    for contract, sources in groups.items():
        failures_before = audit.report.failures
        audit.log(f'合同 {contract}：先检查三份文件是否齐全', 'TITLE')
        for kind in KINDS:
            paths = sources.get(kind, [])
            audit.check(len(paths) == 1, f'{contract} / {kind} 文件',
                        '、'.join(p.name for p in paths) if paths else f'缺少 {kind}{contract}.xls 或 .xlsx')
            if len(paths) > 1:
                audit.log('同一合同同类文件重复，请移出多余文件后重新选择目录；本次不自动挑选。', 'ERROR')
        documents = audit.report.documents[contract] = {}
        for kind in KINDS:
            paths = sources.get(kind, [])
            if len(paths) != 1:
                continue
            try:
                documents[kind] = read_document(paths[0], kind, contract, audit)
            except Exception as exc:
                audit.check(False, f'{paths[0].name} 读取', str(exc))
        compare_documents(documents, contract, audit)
        audit.report.contract_failures[contract] = audit.report.failures - failures_before
    result = audit.report
    audit.log(f'检查结束：{len(groups)} 个合同，{result.checks} 项检查，{result.failures} 项未通过。'
              + ('三份单据所核对的信息一致。' if result.passed else '存在不一致或无法核对的资料，请查看红色日志。'),
              'SUCCESS' if result.passed else 'ERROR')
    return result


def result_messages(report):
    """User-facing results only; the full trace remains in report.entries and the daily file."""
    messages = []
    for contract in report.documents:
        if report.contract_failures.get(contract, 0) == 0:
            messages.append((f'合同 {contract}：核对通过，三份单据齐全，合同号、品名、数量、单价和金额一致。', 'SUCCESS'))
        else:
            messages.append((f'合同 {contract}：核对未通过。', 'ERROR'))
    for title, details in report.issues:
        # Keep the file and differing values, but leave spreadsheet coordinates in the detailed log.
        details = re.sub(r'\s*\([^()]*![A-Z]+\d+\)', '', details)
        details = details.replace('详见上方各单据数值', '详情见日志文件')
        messages.append((f'{title}：{details}', 'ERROR'))
    return messages
