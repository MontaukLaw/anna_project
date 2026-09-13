"""订单排期工作簿的后台预读取，保留每张工作表的原始行列。"""
from dataclasses import dataclass
from pathlib import Path
from queue import Queue
from openpyxl import load_workbook
from openpyxl.utils.datetime import WINDOWS_EPOCH
from services.settings_service import save_file_path


@dataclass
class OrderSchedule:
    source: Path
    sheets: dict[str, list[tuple]]
    epoch: object = WINDOWS_EPOCH

    @property
    def nonempty_rows(self):
        return sum(any(value is not None for value in row)
                   for rows in self.sheets.values() for row in rows)


def read_order_schedule(path: Path) -> OrderSchedule:
    if path.suffix.lower() != ".xlsx":
        raise ValueError("请选择 .xlsx 格式的订单排期表")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook}
    finally:
        workbook.close()
    result = OrderSchedule(path, sheets, workbook.epoch)
    if not result.nonempty_rows:
        raise ValueError("订单排期表没有可读取的内容")
    return result


def load_order_schedule_job(path: Path, config_path: Path, remember: bool, events: Queue):
    try:
        schedule = read_order_schedule(path)
        events.put(("schedule_loaded", schedule))
        if remember:
            try:
                save_file_path(config_path, path, "order_schedule")
            except Exception as exc:
                events.put(("schedule_config_error", str(exc)))
            else:
                events.put(("schedule_remembered", str(config_path)))
    except Exception as exc:
        events.put(("schedule_error", str(exc)))
    finally:
        events.put(("schedule_done", None))
